"""验证水印位置偏移（上下左右独立调整）生效。"""
import os, tempfile
from io import BytesIO
from docx import Document
from docx.oxml.ns import qn

from PIL import Image

from watermark_tool import engine_docx, preview


def count_marks(path):
    doc = Document(path)
    t = sum(1 for s in doc.sections for h in _headers(doc, s)
            for d in h._element.iter(qn("w:drawing"))
            if (d.find(".//" + qn("wp:docPr")) is not None
                and d.find(".//" + qn("wp:docPr")).get("name") == engine_docx.MARK_TEXT))
    i = sum(1 for s in doc.sections for h in _headers(doc, s)
            for d in h._element.iter(qn("w:drawing"))
            if (d.find(".//" + qn("wp:docPr")) is not None
                and d.find(".//" + qn("wp:docPr")).get("name") == engine_docx.MARK_IMG))
    return t, i


def _headers(doc, section):
    hs = [section.header]
    return hs


def _first_posoffset(path, name):
    doc = Document(path)
    for section in doc.sections:
        for h in _headers(doc, section):
            for d in h._element.iter(qn("w:drawing")):
                dp = d.find(".//" + qn("wp:docPr"))
                if dp is not None and dp.get("name") == name:
                    posH = d.find(".//" + qn("wp:positionH"))
                    posV = d.find(".//" + qn("wp:positionV"))
                    x = int(posH.find(qn("wp:posOffset")).text)
                    y = int(posV.find(qn("wp:posOffset")).text)
                    # anchor 内是否还存在 align 居中（不应存在）
                    has_align = posH.find(qn("wp:align")) is not None
                    return x, y, has_align
    return None


def main():
    tmp = tempfile.mkdtemp()
    src = os.path.join(tmp, "src.docx")
    img = os.path.join(tmp, "wm.png")
    Image.new("RGBA", (200, 200), (0, 128, 255, 255)).save(img, "PNG")
    d = Document()
    d.add_paragraph("正文")
    d.save(src)

    # 文本居中 + 图像右移 40%、下移 30%
    opts = {
        "text": {"text": "机密", "font_size": 120, "color": (100, 100, 100),
                 "angle": 0.0, "transparency": 0.3, "scale": 1.0,
                 "offset_x": 0.0, "offset_y": 0.0},
        "image": {"image_path": img, "angle": 0.0, "transparency": 0.6, "scale": 0.7,
                  "offset_x": 40.0, "offset_y": 30.0},
    }
    out = os.path.join(tmp, "out.docx")
    import shutil
    shutil.copy2(src, out)  # core 实际会先把源文件复制到输出，再在副本上加水印
    res = engine_docx.insert_watermark(out, ["text", "image"], **opts)
    print("insert:", res)

    t, i = count_marks(out)
    assert t >= 1 and i >= 1, f"marks missing t={t} i={i}"

    tx, ty, talign = _first_posoffset(out, engine_docx.MARK_TEXT)
    ix, iy, ialign = _first_posoffset(out, engine_docx.MARK_IMG)
    print(f"text pos x={tx} y={ty} align={talign}")
    print(f"image pos x={ix} y={iy} align={ialign}")
    assert not talign and not ialign, "仍存在居中 align，应改用 posOffset 绝对定位"
    assert ix > tx, f"图像应比文本更靠右：ix={ix} tx={tx}"
    assert iy > ty, f"图像应比文本更靠下：iy={iy} ty={ty}"

    # 预览：偏移应改变绘制位置（右上角 vs 居中）
    img_centered = preview.render_preview(
        {"text": {"text": "机密", "font_size": 120, "color": (100, 100, 100),
                  "angle": 0.0, "transparency": 0.3, "scale": 1.0,
                  "offset_x": 0.0, "offset_y": 0.0}},
        ["text"])
    img_shifted = preview.render_preview(
        {"text": {"text": "机密", "font_size": 120, "color": (100, 100, 100),
                  "angle": 0.0, "transparency": 0.3, "scale": 1.0,
                  "offset_x": 40.0, "offset_y": -30.0}},
        ["text"])
    # 找非白像素的包围盒中心，验证发生了位移
    def bbox_center(im):
        import numpy as np
        a = np.array(im.convert("RGB"))
        mask = (a < 240).all(axis=2)
        ys, xs = np.where(mask)
        return (xs.mean(), ys.mean())
    c0 = bbox_center(img_centered)
    c1 = bbox_center(img_shifted)
    print(f"preview center no-offset={c0} offset={c1}")
    assert c1[0] > c0[0] + 10, "预览水平偏移未生效"
    assert c1[1] < c0[1] - 10, "预览垂直偏移未生效"

    print("OK: 文本/图像上下左右偏移均生效，且各自独立")


if __name__ == "__main__":
    main()
