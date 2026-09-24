import os, sys
from docx import Document
sys.path.insert(0, os.path.dirname(__file__))
from watermark_tool import engine_docx as eng

OUT = os.path.join(os.path.dirname(__file__), "sample.docx")

def make_sample():
    d = Document()
    d.add_heading("测试文档", level=1)
    d.add_paragraph("这是正文内容，水印不应影响它。")
    # 第二节能触发多节页眉逻辑
    d.add_section()
    d.add_paragraph("第二节的正文。")
    d.save(OUT)
    print("样例已生成:", OUT)

def main():
    make_sample()
    # 文本水印
    r = eng.insert_watermark(OUT, "text", text="机密 CONFIDENTIAL",
                             color=(200, 0, 0), angle=45, transparency=0.6, scale=1.0)
    print("插入文本水印:", r)
    print("检测水印存在:", eng.has_watermark(OUT))

    # 重复插入不应叠加
    r2 = eng.insert_watermark(OUT, "text", text="机密 CONFIDENTIAL", transparency=0.6)
    print("重复插入:", r2, "（inserted 计数应接近上一轮）")

    # 清除
    rc = eng.clear_watermark(OUT)
    print("清除:", rc)
    print("清除后检测水印存在:", eng.has_watermark(OUT))

    # 图片水印（用 Pillow 现造一张）
    from PIL import Image
    tmpimg = os.path.join(os.path.dirname(__file__), "wm_logo.png")
    Image.new("RGBA", (300, 150), (0, 120, 200, 255)).save(tmpimg)
    ri = eng.insert_watermark(OUT, "image", image_path=tmpimg, angle=30, transparency=0.7, scale=1.0)
    print("插入图片水印:", ri)
    print("检测图片水印存在:", eng.has_watermark(OUT))
    rci = eng.clear_watermark(OUT)
    print("清除图片水印:", rci, "剩余:", eng.has_watermark(OUT))

    # 验证正文未被破坏
    d2 = Document(OUT)
    paras = [p.text for p in d2.paragraphs if p.text.strip()]
    print("正文段落保留:", paras)

if __name__ == "__main__":
    main()
