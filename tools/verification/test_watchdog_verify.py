"""
后台守护（水印被删自动补回）端到端验证脚本。

无需打开 GUI、无需手工操作，直接运行：
    python test_watchdog_verify.py

覆盖三种“水印消失”的场景：
  场景A：被人用工具清除（等价于一键清除）
  场景B：有人在 Word 里打开页眉、手动删掉水印图形
  场景C：输出文件整个被删除
每种场景后等待守护的一个检查周期，验证水印/文件是否被自动补回。
"""
import os
import sys
import time

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.oxml.ns import qn


sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
from watermark_tool import core, engine_docx, watchdog

TMP = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_wd_verify_tmp")
SRC = os.path.join(TMP, "源文件.docx")
OUT = os.path.join(TMP, "输出_水印.docx")

OPTS = {
    "text": {
        "text": "机密 TEST",
        "cn_font_name": "微软雅黑",
        "latin_font_name": "Times New Roman",
        "font_size": 120,
        "color": (128, 128, 128),
        "angle": 45,
        "transparency": 0.5,
        "scale": 1.0,
        "offset_x": 0,
        "offset_y": 0,
    },
    "image": {},
}
KINDS = ["text"]
INTERVAL = 0.5          # 守护检查间隔（秒），与 GUI 里“检查间隔”一致
WAIT = INTERVAL * 3     # 等待足够一个周期完成补回


def make_src():
    os.makedirs(TMP, exist_ok=True)
    if os.path.exists(OUT):
        os.remove(OUT)
    d = Document()
    d.add_paragraph("正文：这是一段用于验证后台守护的测试内容。")
    for i in range(3):
        d.add_paragraph(f"测试段落 {i + 1}")
    d.save(SRC)


def manual_remove_in_word(path):
    """模拟“某人在 Word 里打开页眉、手动删除水印图形”的效果（等价的 XML 操作）。"""
    doc = Document(path)
    removed = 0
    for section in doc.sections:
        for header in (section.header,):
            for drawing in list(header._element.iter(qn("w:drawing"))):
                docPr = drawing.find(".//" + qn("wp:docPr"))
                if docPr is not None and docPr.get("name") in engine_docx.MARK_NAMES:
                    run = drawing.getparent()
                    if run is not None and run.getparent() is not None:
                        run.getparent().remove(run)
                        removed += 1
    doc.save(path)
    return removed


def safe_has_watermark(path):
    """守护重建文件途中去读会抛 PackageNotFoundError（文件尚未写完），按“无水印”处理。"""
    try:
        return core.has_watermark(path)
    except Exception:
        return False


def wait_for(predicate, timeout=WAIT, label=""):
    """轮询等待条件成立（比固定 sleep 稳，避免机器慢时误判失败）。"""
    end = time.time() + timeout
    while time.time() < end:
        if predicate():
            return True
        time.sleep(0.1)
    return False


def main():
    print("=" * 60)
    print("后台守护验证：水印被删后是否自动补回")
    print("=" * 60)
    make_src()

    # 准备：先在输出文件上插入水印
    core.insert_watermark(SRC, KINDS, output_path=OUT, **OPTS)
    assert core.has_watermark(OUT), "准备失败：插入后输出文件应带有水印"
    print(f"[准备] 已生成带水印的输出文件: {OUT}\n")

    wd = watchdog.WatermarkWatchdog(SRC, OUT, KINDS, OPTS, interval=INTERVAL,
                                    log=lambda m, verbose=False: print("   [守护]", m))
    wd.start()
    time.sleep(INTERVAL * 2)
    ok = True

    # 场景A：被“一键清除”掉
    print("\n【场景A】模拟他人用工具把水印清除")
    engine_docx.clear_watermark(OUT)
    gone = not safe_has_watermark(OUT)
    print(f"   清除后水印存在? {not gone}（应为 False）")
    back = wait_for(lambda: safe_has_watermark(OUT))
    print(f"   → 守护自动补回? {back}")
    ok &= gone and back

    # 场景B：在 Word 里手动删除页眉中的水印图形
    print("\n【场景B】模拟有人在 Word 里手动删除页眉水印图形")
    n = manual_remove_in_word(OUT)
    gone = (n > 0) and (not safe_has_watermark(OUT))
    print(f"   手动删除 {n} 个水印图形，删除后水印存在? {not gone}（应为 False）")
    back = wait_for(lambda: safe_has_watermark(OUT))
    print(f"   → 守护自动补回? {back}")
    ok &= bool(gone) and back

    # 场景C：输出文件整个被删除
    print("\n【场景C】模拟输出文件整个被删除")
    os.remove(OUT)
    print(f"   已删除输出文件，存在? {os.path.exists(OUT)}（应为 False）")
    back = wait_for(lambda: os.path.exists(OUT) and safe_has_watermark(OUT))
    print(f"   → 守护自动重建文件并补回水印? {back}")
    ok &= back

    wd.stop()
    print("\n" + "=" * 60)
    print("结论:", "守护功能正常，三种场景均能自动补回 ✅" if ok else
          "存在未通过的场景 ❌（见上方各场景的 → 行）")
    print("=" * 60)
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
