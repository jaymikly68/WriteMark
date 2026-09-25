"""测试：多节水印（每节页眉各写一份）+ 预览渲染。

注意：本文件曾因 API 变更而过期（旧写法把 text/image 参数平铺传入），
现按当前接口修正：参数按 kinds=["text"] + opts={"text": {...}} 传入。
"""
import os
import sys
import zipfile

from docx import Document

sys.path.insert(0, ".")

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
from watermark_tool import engine_docx, preview

MS = "test_multisec.docx"
if os.path.exists(MS):
    os.remove(MS)

doc = Document()
doc.add_paragraph("第一节正文。")
sec2 = doc.add_section()  # 新增一节
# 强制第二节使用独立（非链接）页眉，触发第二个 header part
sec2.header.is_linked_to_previous = False
doc.add_paragraph("第二节正文。")
doc.save(MS)

TEXT = "机密 CONFIDENTIAL"
opts = {
    "text": {
        "text": TEXT, "font_size": 120, "color": (128, 128, 128),
        "angle": 45, "transparency": 0.5, "scale": 1.0,
    }
}

res = engine_docx.insert_watermark(MS, ["text"], **opts)
print("insert:", res)

# 检查每个 header 的图片关系是否都拿到了真实图片数据
with zipfile.ZipFile(MS) as z:
    rels = [n for n in z.namelist() if "header" in n and n.endswith(".rels")]
    media = [n for n in z.namelist() if n.startswith("word/media/")]
    print("header rels:", rels)
    print("media files:", media)
    for m in media:
        data = z.read(m)
        print(f"  {m}: {len(data)} bytes")
    assert media, "没有任何水印图片被写入"

# 预览
img = preview.render_preview(opts, ["text"])
img.save("preview_sample.png")
print("preview size:", img.size)
print("OK")
