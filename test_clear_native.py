"""验证「一键清除」能去掉原本就带水印（非本工具添加）的 Word。

模拟一个 Word 原生水印：页眉里有一个 behindDoc=1 的 drawing，但不带本工具标记名。
clear_watermark 应把它也删掉。
"""
import os, tempfile, shutil
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from watermark_tool import engine_docx

# 一个最简的、衬于文字下方的 drawing（模拟 Word 原生水印）
_NATIVE_DRAWING = (
    '<w:drawing xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
    '<wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="0" '
    'behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">'
    '<wp:simplePos x="0" y="0"/>'
    '<wp:positionH relativeFrom="page"><wp:align>center</wp:align></wp:positionH>'
    '<wp:positionV relativeFrom="page"><wp:align>center</wp:align></wp:positionV>'
    '<wp:extent cx="1000000" cy="500000"/>'
    '<wp:effectExtent l="0" t="0" r="0" b="0"/>'
    '<wp:wrapNone/>'
    '<wp:docPr id="999" name="WordPictureWatermark123" descr=""/>'
    '<wp:cNvGraphicFramePr/>'
    '<a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
    '<a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
    '<pic:pic xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture">'
    '<pic:nvPicPr><pic:cNvPr id="0" name="x"/><pic:cNvPicPr/></pic:nvPicPr>'
    '<pic:blipFill><a:blip r:embed="rId99"/>'
    '<a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
    '<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="1000000" cy="500000"/></a:xfrm>'
    '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr>'
    '</pic:pic></a:graphicData></a:graphic>'
    '</wp:anchor></w:drawing>'
)


def _inject_native_watermark(path):
    doc = Document(path)
    header = doc.sections[0].header
    if not header.paragraphs:
        header.add_paragraph()
    p = header.paragraphs[0]
    r = p.add_run()
    drawing = etree.fromstring(_NATIVE_DRAWING.encode("utf-8"))
    r._r.append(drawing)
    # 注册命名空间，避免序列化前缀异常
    for pref, uri in (("a", "http://schemas.openxmlformats.org/drawingml/2006/main"),
                      ("pic", "http://schemas.openxmlformats.org/drawingml/2006/picture"),
                      ("wp", "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"),
                      ("r", "http://schemas.openxmlformats.org/officeDocument/2006/relationships")):
        etree.register_namespace(pref, uri)
    doc.save(path)


def _count_drawings(path):
    doc = Document(path)
    n = 0
    for section in doc.sections:
        for d in section.header._element.iter(qn("w:drawing")):
            n += 1
    return n


def main():
    tmp = tempfile.mkdtemp()
    src = os.path.join(tmp, "native_wm.docx")
    d = Document()
    d.add_paragraph("正文内容")
    d.save(src)

    # 注入一个非本工具标记的原生水印
    _inject_native_watermark(src)
    assert _count_drawings(src) == 1, "注入原生水印失败"
    print("注入后 drawing 数:", _count_drawings(src))

    # 一键清除（save-as 模型：clear 作用在输出副本，这里直接对 src 做等价验证）
    out = os.path.join(tmp, "cleared.docx")
    shutil.copy2(src, out)
    res = engine_docx.clear_watermark(out)
    print("clear 结果:", res)

    remaining = _count_drawings(out)
    print("清除后 drawing 数:", remaining)
    assert remaining == 0, f"原生水印未被清除，剩余 {remaining} 个 drawing"
    print("OK: 一键清除可去掉原本带水印（非本工具添加）的 Word")


if __name__ == "__main__":
    main()
