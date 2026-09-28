"""
验证「清除水印」的原生水印守门（gui.App._clear）：
- 无 WriteMark 水印、但有原生水印 -> 必须弹窗确认，用户拒绝则完全不执行清除；
- 用户确认后才执行（kinds=None 走全清，含原生水印）；
- 既无工具水印也无原生水印 -> 告知无内容可清，不执行；
- 有本工具水印 -> 按类型精确清除，不触发原生确认；
  （v1.6.3 起）若同时存在 Word/用户水印，会先走「四选一冲突框」`_ask_clear_conflict`。
"""
from __future__ import annotations

import os
import sys
import tempfile
from unittest import mock

import pytest
from docx import Document
from PySide6.QtWidgets import QApplication

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))

from watermark_tool import engine_docx
from watermark_tool.gui import App


def _make_docx():
    tmp = tempfile.mkdtemp()
    p = os.path.join(tmp, "d.docx")
    Document().save(p)
    return p


def _clear_with(w, monkeypatch, types, native, accept=None):
    """在打桩（types/native）下调用一次 _clear，返回 (run_called, info_calls)。"""
    calls = []
    info_calls = []
    monkeypatch.setattr(w, "_run", lambda fn: calls.append(fn))
    monkeypatch.setattr(engine_docx, "detect_watermark_types", lambda path: types)
    monkeypatch.setattr(engine_docx, "detect_native_watermark", lambda path: native)
    monkeypatch.setattr("watermark_tool.gui.QMessageBox.information",
                        lambda *a, **k: info_calls.append(k))
    if accept is not None:
        monkeypatch.setattr(w, "_ask_clear_native", lambda: accept)
    w._clear()
    return bool(calls), info_calls


def test_native_guard_declined_does_not_clear(app, monkeypatch):
    w = App()
    w.file_path = _make_docx()
    w.out_edit.setText("")
    ran, _ = _clear_with(w, monkeypatch, types=set(), native=True, accept=False)
    assert ran is False, "用户拒绝清除原生水印时，不应执行任何清除动作"


def test_native_guard_accepted_runs_clear(app, monkeypatch):
    w = App()
    w.file_path = _make_docx()
    w.out_edit.setText("")
    ran, _ = _clear_with(w, monkeypatch, types=set(), native=True, accept=True)
    assert ran is True, "用户确认后应执行一次清除"


def test_no_watermark_at_all_informs_and_skips(app, monkeypatch):
    w = App()
    w.file_path = _make_docx()
    w.out_edit.setText("")
    ran, info_calls = _clear_with(w, monkeypatch, types=set(), native=False)
    assert ran is False, "无任何水印时应跳过清除"
    assert info_calls, "应提示『未检测到任何水印』"


def test_tool_watermark_no_native_prompt(app, monkeypatch):
    """有本工具水印时不应弹出「仅清原生水印」那个守门框（_ask_clear_native）。

    v1.6.3 起的新行为：本工具水印 与 Word/用户水印**同时存在**时，会先弹出
    「四选一冲突框」`_ask_clear_conflict`（本工具/用户/两者/取消）。
    这条用例打桩选「仅清除本工具水印」——因为要断言的对象是**另一个**守门
    `_ask_clear_native`（无工具水印时的全清确认）从未被调用。
    不打桩的话 `QMessageBox.exec()` 会真的弹出模态框把测试挂住。
    """
    w = App()
    w.file_path = _make_docx()
    w.out_edit.setText("")
    prompts = []
    monkeypatch.setattr(w, "_ask_clear_conflict", lambda: "tool")
    ran, _ = _clear_with(w, monkeypatch, types={"text"}, native=True)
    assert ran is True
    # 再单独打桩 _ask_clear_native，重跑一次（types={'text'} 分支不会走到原生确认）
    monkeypatch.setattr(w, "_ask_clear_native", lambda: prompts.append(1) or True)
    w._clear()
    assert prompts == [], "有本工具水印时不应弹出原生水印确认"


# ---------------------------------------------------------------------------
# 端到端（不打桩清除本身）：真实文件 + 真实守门 + 真实 core.clear_watermark
# 回归背景：用户的手工文本框水印（AlternateContent + 旋转文本框）曾整链路失效
# （不弹窗 + removed=0 + 输出文件水印原样）。
# ---------------------------------------------------------------------------

_WPS_URI = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
_NS = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:v="urn:schemas-microsoft-com:vml" '
    'xmlns:o="urn:schemas-microsoft-com:office:office"'
)
_AC_TAG = "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"


def _make_user_textbox_wm_docx():
    """构造与用户真实文档同构的水印文档：页眉里旋转文本框水印（含 VML Fallback）。"""
    from docx.oxml import parse_xml
    tmp = tempfile.mkdtemp()
    p = os.path.join(tmp, "水印.docx")
    doc = Document()
    doc.add_paragraph("正文内容一行")
    ac = parse_xml(
        f'<mc:AlternateContent {_NS}>'
        '<mc:Choice Requires="wps">'
        '<w:drawing><wp:anchor behindDoc="0" distT="0" distB="0" distL="0" distR="0" '
        'simplePos="0" relativeHeight="1" locked="0" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="margin"><wp:posOffset>0</wp:posOffset></wp:positionH>'
        '<wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV>'
        '<wp:extent cx="100000" cy="100000"/><wp:wrapNone/>'
        '<wp:docPr id="903" name="文本框 2"/>'
        f'<a:graphic><a:graphicData uri="{_WPS_URI}">'
        '<wps:wsp><wps:cNvSpPr txBox="1"/><wps:spPr>'
        '<a:xfrm rot="18900000"><a:off x="0" y="0"/><a:ext cx="100000" cy="100000"/></a:xfrm>'
        '</wps:spPr></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing>'
        '</mc:Choice>'
        '<mc:Fallback><w:pict>'
        '<v:shape id="文本框 2" style="position:absolute;rotation:-45;z-index:251659264">'
        '<v:textbox><w:txbxContent><w:r><w:t>丽珠医药</w:t></w:r></w:txbxContent></v:textbox>'
        '</v:shape></w:pict></mc:Fallback>'
        '</mc:AlternateContent>'
    )
    hdr = doc.sections[0].header
    (hdr.paragraphs[0] if hdr.paragraphs else hdr.add_paragraph()).add_run()._r.append(ac)
    doc.save(p)
    return p


def test_e2e_user_watermark_prompt_then_real_clear(app, monkeypatch):
    """端到端：用户自己的文本框水印 -> 弹窗(模拟点“是”) -> 输出文件真的没水印。"""
    from watermark_tool import core
    src = _make_user_textbox_wm_docx()
    w = App()
    w.file_path = src
    out = os.path.join(tempfile.mkdtemp(), "水印WaterMark.docx")
    w.out_edit.setText(out)
    # 守门必须弹（真实检测，不打桩 detect）；模拟用户点“是”
    monkeypatch.setattr(w, "_ask_clear_native", lambda: True)
    jobs = []
    monkeypatch.setattr(w, "_run", lambda fn: jobs.append(fn))
    w._clear()
    assert jobs, "守门确认后必须执行清除"
    r = jobs[0]()  # 真正跑 core.clear_watermark
    assert r["ok"] and r["removed"] >= 1, f"输出文件应被真实清除，实际 {r}"
    doc2 = Document(out)
    assert "正文内容一行" in "".join(par.text for par in doc2.paragraphs), "正文必须完好"
    hdr_el = doc2.sections[0].header._element
    assert not list(hdr_el.iter(_AC_TAG)), "AlternateContent 水印块必须整块删除（含 Fallback）"
    assert engine_docx.detect_native_watermark(out) is False, "清除后不应再检测到水印"


def test_e2e_user_watermark_declined_no_output_change(app, monkeypatch):
    """端到端反向：用户点“否” -> 完全不执行，输出路径不会产生文件。"""
    src = _make_user_textbox_wm_docx()
    w = App()
    w.file_path = src
    out = os.path.join(tempfile.mkdtemp(), "水印WaterMark.docx")
    w.out_edit.setText(out)
    monkeypatch.setattr(w, "_ask_clear_native", lambda: False)
    ran, _ = _clear_with(w, monkeypatch, types=set(), native=True, accept=False)
    assert ran is False
    assert not os.path.exists(out), "用户拒绝后不应生成输出文件"


if __name__ == "__main__":
    app = QApplication([])
    pytest.main([__file__, "-q"])
    app.quit()
