"""
「插入/清除进行中关闭窗口」——走 GUI 真实 `_run()` 路径的回归测试。

与 tests/gui/test_close.py 的区别：那边直接 `App._worker = Worker(长任务)` 绕过
调度；这里从 `App._run()` 进入，因此**同时验证 UI 忙碌态**（按钮禁用、状态栏
"处理中…"）与关闭时的线程终止，覆盖真实用户点「一键插入水印」后立刻关窗的时序。

如何制造“任务仍在跑”：
    用 `Worker.pre_hook`（**仅测试注入**的可控阻塞点，生产路径恒为 None）。
    刻意**不给正常插入流程加任何 sleep**——用户实测 30 页插入一两秒就完成，
    那是正常速度，不是缺陷；测试只能自己造阻塞点，不能拖慢真实用户。
"""
import os
import sys
import tempfile
import threading
import time

sys.path.insert(0, '.')

from PySide6.QtWidgets import QApplication, QLineEdit
from PySide6.QtGui import QCloseEvent
from docx import Document

from watermark_tool import core
from watermark_tool.gui import App, Worker


class _Gate:
    """测试专用阻塞点：**只能**用 time.sleep 轮询，不能用 threading.Event.wait。

    实测：线程阻塞在 `Event.wait()`（Windows 的 WaitForSingleObject）时被
    `QThread.terminate()` 强杀，会让整个 pytest 进程 access violation 崩溃。
    time.sleep 轮询则可以被安全终止。closeEvent 的兜底是 terminate，
    所以阻塞点必须以“可被 terminate 安全打断”的方式实现。
    """

    def __init__(self, seconds=10.0):
        self._seconds = seconds
        self._open = True

    def block(self):
        end = time.time() + self._seconds
        while self._open and time.time() < end:
            time.sleep(0.05)

    def release(self):
        self._open = False


def test_production_worker_has_no_hook():
    """生产路径不得带任何测试 hook（保证真实插入速度不被拖慢）。"""
    w = Worker(lambda: None)
    assert w.pre_hook is None, "正常 Worker 不应带 pre_hook"


def test_run_path_busy_state(app):
    """任务运行期间 GUI 必须处于忙碌态：按钮禁用、状态栏提示处理中。"""
    w = App()
    w.show()
    gate = _Gate()
    w._test_worker_hook = gate.block                  # 测试专用阻塞点

    d = tempfile.mkdtemp()
    src = os.path.join(d, "in.docx")
    doc = Document()
    doc.add_paragraph("正文")
    doc.save(src)
    out = os.path.join(d, "out.docx")
    w.file_path, w.out_edit = src, QLineEdit(out)

    w._run(lambda: core.insert_watermark(src, ["text"], output_path=out,
                                         text={"text": "机密"}))
    try:
        assert w._worker is not None and w._worker.isRunning(), "任务应已启动"
        assert not w._btn_insert.isEnabled(), "处理期间必须禁用「一键插入」防止重入"
        assert not w._btn_clear.isEnabled(), "处理期间必须禁用「一键清除」防止重入"
        assert "处理中" in w.status_label.text(), (
            f"状态栏应显示处理中，实际：{w.status_label.text()}")
        print("[gui] 插入进行中：按钮禁用 + 状态栏「处理中…」")
    finally:
        gate.release()
        if w._worker is not None:
            w._worker.wait(3000)


def test_close_during_insert_terminates_and_leaves_no_partial_output(app):
    """插入被阻塞时关闭：线程必须被终止，且**不得留下半截输出文件**。"""
    w = App()
    w.show()
    gate = _Gate()
    w._test_worker_hook = gate.block                  # 卡在任务真正执行之前

    d = tempfile.mkdtemp()
    src = os.path.join(d, "in.docx")
    doc = Document()
    doc.add_paragraph("正文")
    doc.save(src)
    out = os.path.join(d, "out.docx")
    w.file_path, w.out_edit = src, QLineEdit(out)

    w._run(lambda: core.insert_watermark(src, ["text"], output_path=out,
                                         text={"text": "机密"}))
    # 等一小会儿，确保线程已进入 pre_hook（生产环境里这一步是真正的插入）
    time.sleep(0.2)
    worker = w._worker
    assert worker is not None and worker.isRunning(), "前置条件：任务应仍在运行"

    t0 = time.time()
    w.closeEvent(QCloseEvent())
    dt = time.time() - t0
    print(f"[gui] 插入进行中关闭 -> closeEvent 耗时 {dt*1000:.0f} ms")
    assert dt < 4.0, f"关闭被卡住（{dt:.1f}s），工作线程未及时终止"
    # 注意：closeEvent 内 wait() 会派发 finished -> _on_worker_finished 把引用清成
    # None（这正是期望的清理行为），所以这里必须用关闭前保存下来的对象来断言
    assert not worker.isRunning(), "工作线程仍 running，退出会卡住"
    assert w._worker is None, "关闭后应释放 worker 引用"
    assert not os.path.exists(out), (
        "任务被终止时不应产生输出文件——半截文件比没有文件更糟")

    gate.release()
    print("[gui] 关闭即终止插入：无卡顿、无半截输出文件")


if __name__ == "__main__":
    app = QApplication([])
    test_production_worker_has_no_hook()
    test_run_path_busy_state(app)
    test_close_during_insert_terminates_and_leaves_no_partial_output(app)
    print("\nALL PASS：插入中关闭窗口行为正确")
