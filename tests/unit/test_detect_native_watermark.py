"""
验证「原生水印检测」(engine_docx.detect_native_watermark)：
它是 _clear 里“确认后才删原生水印”守门的依据——命中条件必须与
clear_watermark(kinds=None) 的实际删除范围保持一致（不带本工具标记的
水印类图形块：behindDoc 浮动图 / 旋转浮动图形 / VML textpath 等），
否则守门就会放跑本该问用户的删除、或把用户正常内容误判成水印。

回归背景（真实用户文档踩坑）：用户手工在页眉平铺的“文本框水印”
（wordprocessingShape + 旋转 + AlternateContent 包裹）曾被完全漏检、
清除后 removed=0；且只删 mc:Choice 里的 drawing 时 mc:Fallback 的
VML 副本仍会渲染。本文件同时固化这三个教训。
"""
from __future__ import annotations

import io
import os
import sys
import tempfile

from docx import Document
from docx.oxml import parse_xml
from PIL import Image

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from watermark_tool import engine_docx


def _png(color=(5, 5, 5), size=(32, 32)):
    b = io.BytesIO()
    Image.new("RGB", size, color).save(b, "PNG")
    return b.getvalue()


def _make_doc(path):
    doc = Document()
    doc.add_paragraph("正文第一段")
    doc.save(path)
    return path


def _native_watermark(hdr):
    """模拟 Word 原生水印：behindDoc=1、无本工具标记的图片。"""
    rId, _ = hdr.part.get_or_add_image(io.BytesIO(_png()))
    return engine_docx._make_drawing(rId, 400000, 400000, 0,
                                    "PowerPlusWaterMarkObject", "", 901, 0, 0)


def _user_behind_picture(hdr):
    """用户自己的衬于文字下方的图片（无本工具标记）—— 与原生水印同构。"""
    rId, _ = hdr.part.get_or_add_image(io.BytesIO(_png()))
    return engine_docx._make_drawing(rId, 400000, 400000, 0,
                                    "用户页眉背景图", "", 902, 0, 0)


_NS = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:v="urn:schemas-microsoft-com:vml" '
    'xmlns:o="urn:schemas-microsoft-com:office:office"'
)


def _textbox_shape_drawing():
    """用户自制“文本框水印”的 Choice drawing（旋转浮动文本框，behindDoc=0）。"""
    xml = (
        f'<w:drawing {_NS}>'
        '<wp:anchor behindDoc="0" distT="0" distB="0" distL="114300" distR="114300" '
        'simplePos="0" relativeHeight="251659264" locked="0" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="margin"><wp:posOffset>1917700</wp:posOffset></wp:positionH>'
        '<wp:positionV relativeFrom="paragraph"><wp:posOffset>189865</wp:posOffset></wp:positionV>'
        '<wp:extent cx="1296000" cy="496800"/><wp:wrapNone/>'
        '<wp:docPr id="878790655" name="文本框 2"/>'
        '<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        '<wps:wsp><wps:cNvSpPr txBox="1"/><wps:spPr>'
        '<a:xfrm rot="18900000"><a:off x="0" y="0"/><a:ext cx="1296000" cy="496800"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></wps:spPr>'
        '<wps:txbx><w:txbxContent><w:p><w:r><w:t>丽珠医药</w:t></w:r></w:p></w:txbxContent></wps:txbx>'
        '</wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing>'
    )
    return parse_xml(xml)


def _ac_block_with_fallback():
    """完整 mc:AlternateContent 块：Choice=DrawingML 文本框 + Fallback=VML。"""
    xml = (
        f'<mc:AlternateContent {_NS}>'
        '<mc:Choice Requires="wps">'
        + _textbox_shape_drawing().getroottree().getroot().tag and ""  # placeholder
    )
    # 上面拼接过于绕，直接一次性写全
    xml = (
        f'<mc:AlternateContent {_NS}>'
        '<mc:Choice Requires="wps">'
        '<w:drawing><wp:anchor behindDoc="0" distT="0" distB="0" distL="0" distR="0" '
        'simplePos="0" relativeHeight="1" locked="0" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="margin"><wp:posOffset>0</wp:posOffset></wp:positionH>'
        '<wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV>'
        '<wp:extent cx="100000" cy="100000"/><wp:wrapNone/>'
        '<wp:docPr id="903" name="文本框 2"/>'
        '<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        '<wps:wsp><wps:cNvSpPr txBox="1"/><wps:spPr>'
        '<a:xfrm rot="18900000"><a:off x="0" y="0"/><a:ext cx="100000" cy="100000"/></a:xfrm>'
        '</wps:spPr></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing>'
        '</mc:Choice>'
        '<mc:Fallback><w:pict>'
        '<v:shape id="文本框 2" o:spid="_x0000_s1026" '
        'style="position:absolute;margin-left:151pt;margin-top:14.95pt;width:102.05pt;'
        'height:39.1pt;rotation:-45;z-index:251659264;mso-position-horizontal:absolute">'
        '<v:textbox><w:txbxContent><w:r><w:t>丽珠医药</w:t></w:r></w:txbxContent></v:textbox>'
        '</v:shape></w:pict></mc:Fallback>'
        '</mc:AlternateContent>'
    )
    return parse_xml(xml)


def _vml_textpath_pict():
    """Word 内置水印的老式裸 w:pict（v:shape + v:textpath）。"""
    xml = (
        f'<w:pict {_NS}>'
        '<v:shape id="PowerPlusWaterMarkObject1" o:spid="_x0000_s2049" '
        'style="position:absolute;margin-left:0;margin-top:0;width:100pt;height:100pt;'
        'z-index:-251654144;mso-position-horizontal:center">'
        '<v:textpath string="机密" style="v-text-align:center"/></v:shape></w:pict>'
    )
    return parse_xml(xml)


def _add_block(hdr, block):
    """把任意顶层图形块追加到页眉第一个段落的新 run 里。"""
    p = hdr.paragraphs[0] if hdr.paragraphs else hdr.add_paragraph()
    r = p.add_run()
    r._r.append(block)


def _add(hdr, drawing):
    engine_docx._add_drawing_to_part(hdr, drawing)


def test_detect_native_true_when_only_native_watermark(tmp_path):
    """文档只有 Word 原生水印（本工具未插过） -> 命中。"""
    p = str(tmp_path / "d.docx")
    _make_doc(p)
    doc = Document(p)
    _add(doc.sections[0].header, _native_watermark(doc.sections[0].header.part))
    doc.save(p)
    assert engine_docx.detect_native_watermark(p) is True


def test_detect_native_false_when_only_tool_watermark(tmp_path):
    """文档只有本工具水印（文本） -> 不命中（那是工具自己的，按类型清、不动原生）。"""
    p = str(tmp_path / "d.docx")
    engine_docx.insert_watermark(_make_doc(p), ["text"], output_path=p)
    assert engine_docx.detect_native_watermark(p) is False


def test_detect_native_false_when_no_watermark(tmp_path):
    """文档没有任何水印 -> 不命中。"""
    p = str(tmp_path / "d.docx")
    _make_doc(p)
    assert engine_docx.detect_native_watermark(p) is False


def test_detect_native_true_when_both_tool_and_native(tmp_path):
    """文档同时有本工具水印与原生水印 -> 原生仍被检测到（守门会按“有原生”处理）。"""
    p = str(tmp_path / "d.docx")
    engine_docx.insert_watermark(_make_doc(p), ["text"], output_path=p)
    doc = Document(p)
    _add(doc.sections[0].header, _native_watermark(doc.sections[0].header.part))
    doc.save(p)
    assert engine_docx.detect_native_watermark(p) is True


def test_detect_native_true_for_user_own_behind_picture(tmp_path):
    """用户在页眉里自己的 behindDoc 图片 -> 判定为原生水印（与 clear_any 删除范围一致，

    守门会对这种“衬于文字下方”的图片弹窗征求同意，而不是静默删除）。
    """
    p = str(tmp_path / "d.docx")
    _make_doc(p)
    doc = Document(p)
    _add(doc.sections[0].header, _user_behind_picture(doc.sections[0].header.part))
    doc.save(p)
    assert engine_docx.detect_native_watermark(p) is True


def test_detect_native_true_for_rotated_textbox_watermark(tmp_path):
    """回归：用户手工的旋转文本框水印（behindDoc=0）必须被检测到。

    真实案例里 11 份水印有 10 份 behindDoc=0，仅靠 behindDoc 会漏检。
    旋转浮动图形是第二判定特征。
    """
    p = str(tmp_path / "d.docx")
    _make_doc(p)
    doc = Document(p)
    _add_block(doc.sections[0].header, _ac_block_with_fallback())
    doc.save(p)
    assert engine_docx.detect_native_watermark(p) is True


def test_detect_native_true_for_vml_textpath(tmp_path):
    """Word 内置水印的老式裸 w:pict（v:textpath）必须被检测到。"""
    p = str(tmp_path / "d.docx")
    _make_doc(p)
    doc = Document(p)
    _add_block(doc.sections[0].header, _vml_textpath_pict())
    doc.save(p)
    assert engine_docx.detect_native_watermark(p) is True


def test_detect_native_false_for_unrotated_inline_textbox(tmp_path):
    """页眉里非浮动、未旋转的普通内容不命中（不误伤）。"""
    p = str(tmp_path / "d.docx")
    _make_doc(p)
    doc = Document(p)
    hdr = doc.sections[0].header
    hdr.paragraphs[0].add_run("页眉普通文字")
    doc.save(p)
    assert engine_docx.detect_native_watermark(p) is False


def test_clear_full_removes_ac_block_including_fallback(tmp_path):
    """全清必须整块删除 AlternateContent：Choice 与 Fallback（VML）都不能残留。"""
    p = str(tmp_path / "d.docx")
    _make_doc(p)
    doc = Document(p)
    _add_block(doc.sections[0].header, _ac_block_with_fallback())
    doc.save(p)
    res = engine_docx.clear_watermark(p)
    assert res["removed"] >= 1
    doc2 = Document(p)
    hdr_el = doc2.sections[0].header._element
    assert not list(hdr_el.iter(engine_docx._AC_TAG)), "AlternateContent 块必须整块删除"
    assert not list(hdr_el.iter(engine_docx._PICT_TAG)), "Fallback 的 w:pict 不能残留"
    assert engine_docx.detect_native_watermark(p) is False


def test_clear_full_removes_textpath_pict(tmp_path):
    """全清要能删掉裸 w:pict 的 v:textpath 内置水印。"""
    p = str(tmp_path / "d.docx")
    _make_doc(p)
    doc = Document(p)
    _add_block(doc.sections[0].header, _vml_textpath_pict())
    doc.save(p)
    res = engine_docx.clear_watermark(p)
    assert res["removed"] >= 1
    doc2 = Document(p)
    hdr_el = doc2.sections[0].header._element
    assert not list(hdr_el.iter(engine_docx._PICT_TAG))
    assert engine_docx.detect_native_watermark(p) is False


def test_clear_kinds_spares_native(tmp_path):
    """按类型清除（kinds=['text']）时，用户的文本框水印必须原样保留。"""
    p = str(tmp_path / "d.docx")
    engine_docx.insert_watermark(_make_doc(p), ["text"], output_path=p)
    doc = Document(p)
    _add_block(doc.sections[0].header, _ac_block_with_fallback())
    doc.save(p)
    res = engine_docx.clear_watermark(p, kinds=["text"])
    assert res["removed"] >= 1
    doc2 = Document(p)
    hdr_el = doc2.sections[0].header._element
    assert list(hdr_el.iter(engine_docx._AC_TAG)), "用户自己的水印不应被按类型清除删掉"
    assert engine_docx.detect_native_watermark(p) is True


if __name__ == "__main__":
    import pytest
    raise SystemExit(pytest.main([__file__, "-q"]))
