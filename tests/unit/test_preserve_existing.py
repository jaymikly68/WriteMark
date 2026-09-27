"""
二次安全审查（commit 84a773e 之后）—— 主题 1：preserve_existing。

要回答的问题：output.docx 被用户改过正文 / 表格 / 图片 / 段落格式之后，
“再次插入水印”这一轮，**用户的修改是否完整保留**。

底层不变式：preserve_existing=True 必须**始终**在现有 output 上操作
（先清后插），而不是在某些分支里悄悄回退成「从 source 重新生成」——
后者会用源文件覆盖掉用户在输出文件上的所有正文修改。

（注：原先挂在这个文件里的“水印被删 → 守护自动补回”用例，依赖已删除的
watermark_tool.watchdog 模块，故已移除；preserve_existing 是独立保留的能力。）
"""
from __future__ import annotations

import io
import os
import sys
import tempfile

from docx import Document
from docx.shared import Pt
from PIL import Image

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from watermark_tool import core, engine_docx


def _png_bytes(color=(10, 20, 30), size=(40, 40)):
    img = Image.new("RGB", size, color)
    bio = io.BytesIO()
    img.save(bio, "PNG")
    return bio.getvalue()


def _make_source(path):
    """造一份「用户原始素材」：正文段落 + 带格式的段落 + 表格 + 正文图片。"""
    doc = Document()
    p = doc.add_paragraph("原始正文第一段")
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(18)
    doc.add_paragraph("原始正文第二段")

    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "表A1"
    t.cell(0, 1).text = "表A2"
    t.cell(1, 0).text = "表B1"
    t.cell(1, 1).text = "表B2"

    doc.add_picture(io.BytesIO(_png_bytes()), width=Pt(120))
    doc.save(path)
    return path


def _text_opts(**over):
    d = {"text": "机密 CONFIDENTIAL", "font_size": 120, "color": (192, 0, 0),
         "angle": 30, "transparency": 0.6, "scale": 1.0,
         "offset_x": 0.0, "offset_y": 0.0}
    d.update(over)
    return {"text": d}


def _snapshot(path):
    """把 output 的“用户可见内容”抓出来，用于补回后逐项比对。

    只记录**非水印**的用户资产：正文文字、字体加粗/字号、表格、正文图片数量。
    水印图形本身不在比对范围内（它本来就期望被重插）。
    """
    doc = Document(path)
    body = []
    for p in doc.paragraphs:
        runs = [(r.text, bool(r.bold), r.font.size.pt if r.font.size else None)
                for r in p.runs]
        body.append((p.text, runs))
    tables = [[[c.text for c in row.cells] for row in t.rows] for t in doc.tables]
    imgs = len(doc.inline_shapes)
    return {"body": body, "tables": tables, "inline_images": imgs}


def test_preserve_existing_keeps_user_body_table_image_format():
    """preserve_existing=True 单独使用：用户资产必须原样留下，只补水印。"""
    tmp = tempfile.mkdtemp()
    src = os.path.join(tmp, "src.docx")
    out = os.path.join(tmp, "out.docx")
    _make_source(src)

    core.insert_watermark(src, ["text"], output_path=out, **_text_opts())
    assert core.watermark_count(out) > 0

    # ---- 用户开始编辑输出文件（守护的目标文件）----
    doc = Document(out)
    doc.paragraphs[0].runs[0].text = "我改过的第一段"
    doc.paragraphs[1].text = "新增的一整段"
    doc.add_paragraph("用户在输出文件里追加的第三段")
    t = doc.tables[0]
    t.cell(0, 0).text = "被用户改成 XYZ"
    t.add_row().cells[0].text = "用户加的行"
    doc.add_picture(io.BytesIO(_png_bytes((200, 30, 30))), width=Pt(60))
    doc.save(out)

    before = _snapshot(out)

    # ---- 用户在已改过的输出文件上再次插入水印（preserve_existing）----
    res = core.insert_watermark(src, ["text"], output_path=out,
                                preserve_existing=True, **_text_opts())
    assert res.get("ok")
    assert core.watermark_count(out) > 0, "插入后水印应存在"

    after = _snapshot(out)
    assert after == before, (
        "preserve_existing 插入后用户资产发生变化——说明发生了覆盖。\n"
        f"before={before}\n\nafter={after}"
    )
    print("[review1] preserve_existing 单独使用：正文/表格/图片/段落格式均完整保留")


def test_preserve_existing_never_reads_source_content():
    """preserve_existing=True 时，output 的内容不应被 source 的内容“反向覆盖”。"""
    tmp = tempfile.mkdtemp()
    src = os.path.join(tmp, "src.docx")
    out = os.path.join(tmp, "out.docx")
    _make_source(src)
    core.insert_watermark(src, ["text"], output_path=out, **_text_opts())

    # 改 source（模拟“源文档后来又改了”），让 output 与 source 明显不同
    doc = Document(src)
    doc.add_paragraph("源文档后来新增的内容（不应被搬进 output）")
    doc.save(src)

    doc = Document(out)
    doc.add_paragraph("output 独有的内容（必须留下）")
    doc.save(out)
    before = _snapshot(out)

    core.insert_watermark(src, ["text"], output_path=out,
                          preserve_existing=True, **_text_opts())

    after = _snapshot(out)
    assert after == before, "output 内容不应随 source 一起变"
    doc2 = Document(out)
    assert "源文档后来新增的内容（不应被搬进 output）" not in [p.text for p in doc2.paragraphs]
    print("[review1] preserve_existing 不回读 source 内容，output 保持独立")


def test_preserve_existing_when_source_is_locked():
    """source 被 Word/WPS 占用时，插入仍应作用于已有 output（此时根本不需要读 source）。"""
    tmp = tempfile.mkdtemp()
    src = os.path.join(tmp, "src.docx")
    out = os.path.join(tmp, "out.docx")
    _make_source(src)
    core.insert_watermark(src, ["text"], output_path=out, **_text_opts())

    # 用户在 output 上改点东西
    doc = Document(out)
    doc.add_paragraph("用户在 output 上的修改")
    doc.save(out)

    # 模拟“其它程序（Word/WPS）持有 source 的独占写句柄”。
    # 注：Python 自带的 open() 在 Windows 上用 _SH_DENYNO，同进程里再开一次不会失败，
    # 因此这里替换占用检测函数，模拟的是「检测结果为被占用」这一事实本身。
    orig = core.is_file_locked
    src_abs = os.path.abspath(src)

    def _fake_locked(path):
        return True if os.path.abspath(path) == src_abs else orig(path)

    core.is_file_locked = _fake_locked
    try:
        engine_docx.clear_watermark(out, kinds=["text"])
        assert core.watermark_count(out) == 0

        # 再次插入：output 存在、且我们只读写 output，应当成功
        core.insert_watermark(src, ["text"], output_path=out,
                              preserve_existing=True, **_text_opts())
        docs = Document(out)
        assert "用户在 output 上的修改" in [p.text for p in docs.paragraphs]
        assert core.watermark_count(out) > 0, (
            "source 被占用时，只要 output 存在就不该连插入一起失败——"
            "preserve_existing 的语义正是“只操作现有 output”。"
        )
    finally:
        core.is_file_locked = orig
    print("[review1] source 被占用时，preserve_existing 仍能插入 output 水印")


if __name__ == "__main__":
    test_preserve_existing_keeps_user_body_table_image_format()
    test_preserve_existing_never_reads_source_content()
    test_preserve_existing_when_source_is_locked()
    print("\nALL PASS：preserve_existing 始终只操作现有 output，用户修改零丢失")
