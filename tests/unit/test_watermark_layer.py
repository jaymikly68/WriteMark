"""
文字水印的**结构层级**回归测试（对应“水印被正文图片遮挡”的修复）。

结论先写在这里，避免后来人重复踩：

    v1.6.1 及更早：水印写在 **页眉 part**，锚点 `wp:anchor` + `behindDoc="1"` +
    `wrapNone` + `relativeFrom="page"`，与 Word 内置水印同一套做法。但 Word 的
    跨层绘制顺序固定为「先页眉/页脚、后正文」——页眉里的图形永远落在正文之下，
    所以正文图片（inline/浮动/behindDoc 任意）都绘制在水印之上，水印经过图片
    区域时会被遮挡。

    v1.6.2 起：水印改放 **正文 body**，锚点 `wp:anchor` + `behindDoc="0"`（浮于文字
    之上）+ `relativeHeight` 取最大值 + `wrapNone` + `relativeFrom="page"`。浮动图形
    锚定到正文的每一个段落（positionV/H 以页面为基准），因此**每一页**都会渲染出
    一份水印、重叠成视觉上的一份，且因为它在正文层、`behindDoc="0"`，必然绘制在
    正文图片之上。

    识别/清除依赖 docPr 上的私有标记（descr=MARK_TEXT/MARK_IMG），与 behindDoc
    无关——层级怎么改都不影响“一键删除”的准确性。

本文件只做**结构层**断言。像素级“水印确实压在图片上”无法在 CI 里验证（需要 Word/WPS
渲染），必须在 Word 页面视图中人工确认；但这里把“浮动于正文之上”的 XML 不变式钉死。
"""
import os
import sys
import tempfile

import pytest
from docx import Document
from PIL import Image

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from watermark_tool import core, engine_docx

qn = engine_docx.qn


def _png(d, name, size=(24, 24), color=(10, 10, 10)):
    p = os.path.join(d, name)
    Image.new("RGB", size, color).save(p)
    return p


def _header_drawings(doc):
    return [dr for s in doc.sections
            for dr in s.header._element.iter(qn("w:drawing"))]


def _body_drawings(doc):
    return list(doc.element.body.iter(qn("w:drawing")))


def _body_marked_drawings(doc):
    return [dr for dr in doc.element.body.iter(qn("w:drawing"))
            if engine_docx._mark_of(dr) is not None]


def _make_doc_with_body_image(d):
    src = os.path.join(d, "body_img.docx")
    doc = Document()
    doc.add_paragraph("用户正文")
    doc.add_picture(_png(d, "bodyimg.png", (60, 40), (200, 200, 200)), width=None)
    doc.save(src)
    return src


def _make_plain_doc(d, name="plain.docx"):
    src = os.path.join(d, name)
    doc = Document()
    doc.add_paragraph("用户正文")
    doc.save(src)
    return src


# ---------------------------------------------------------------------------
def test_watermark_is_placed_in_body_above_content():
    """水印必须落在正文 body（而非页眉）、behindDoc=0、绝对定位、无环绕。

    这几个属性共同构成“浮于正文之上”的语义；v1.6.2 把放位置从页眉换到正文，
    并把 behindDoc 从 1 改为 0，任何一项被改掉都会破坏“水印不被正文图片遮挡”。
    """
    d = tempfile.mkdtemp()
    src = _make_plain_doc(d)          # 纯文本文档：正文里本来就没有图形
    out = os.path.join(d, "out.docx")
    core.insert_watermark(src, ["text"], output_path=out,
                          text={"text": "机密", "font_size": 80})

    doc = Document(out)
    marked = _body_marked_drawings(doc)
    assert len(marked) >= 1, f"正文里应至少有一份水印，实际 {len(marked)}"
    assert _header_drawings(doc) == [], "水印不应被写进页眉（v1.6.2 起在正文）"

    anchor = marked[0].find(".//" + qn("wp:anchor"))
    assert anchor is not None, "水印应是浮动锚定(anchor)图形"
    assert anchor.get("behindDoc") == "0", "水印应浮于文字之上(behindDoc=0)"
    assert marked[0].find(".//" + qn("wp:wrapNone")) is not None, "水印应无环绕"
    pos_h = marked[0].find(".//" + qn("wp:positionH"))
    pos_v = marked[0].find(".//" + qn("wp:positionV"))
    assert pos_h.get("relativeFrom") == "page", "水平定位应以页面为基准"
    assert pos_v.get("relativeFrom") == "page", "垂直定位应以页面为基准"
    print("[层] 水印位于正文 body · behindDoc=0 · relativeFrom=page · wrapNone")


def test_body_image_stays_in_body_layer_and_watermark_floats_above():
    """正文图片留在正文层；水印也进正文层但 behindDoc=0，因此压在图片之上。

    把「遮挡来自跨层绘制顺序」这一事实钉住：旧版水印在页眉、图片在正文，Word 先画
    页眉后画正文，正文图片在上；新版水印与图片同在正文层，但 behindDoc=0 使其浮于
    图片之上（结构层保证；像素级需 Word 人工确认）。
    """
    d = tempfile.mkdtemp()
    src = _make_doc_with_body_image(d)
    out = os.path.join(d, "out.docx")
    core.insert_watermark(src, ["text"], output_path=out,
                          text={"text": "机密", "font_size": 80})

    doc = Document(out)
    body_imgs = [dr for dr in _body_drawings(doc) if engine_docx._mark_of(dr) is None]
    # 正文里的未标记图形（行内图片）必须完好
    assert len(body_imgs) == 1, f"正文图片应留在正文层，实际 {len(body_imgs)}"
    assert len(_body_marked_drawings(doc)) >= 1, "正文里应有浮于其上的水印"
    # 水印必须是浮动且 behindDoc=0（这是“压在图片之上”的结构前提）
    wm = _body_marked_drawings(doc)[0]
    assert wm.find(".//" + qn("wp:anchor")).get("behindDoc") == "0"
    print("[层] 正文图片在 body / 水印也在 body 且 behindDoc=0 —— 浮于图片之上")


def test_watermark_marker_independent_of_behind_doc():
    """识别与清除依赖 docPr 上的私有标记，不依赖 behindDoc。

    v1.6.2 已经把 behindDoc 固定为 "0"；这里再人为改回 "1"（沉到文字下方）验证
    识别/删除逻辑**不需要**跟着改——这个不变式提前钉住，防止将来改动层级时
    悄悄破坏“一键删除”。
    """
    d = tempfile.mkdtemp()
    src = _make_doc_with_body_image(d)
    out = os.path.join(d, "out.docx")
    core.insert_watermark(src, ["text"], output_path=out,
                          text={"text": "机密", "font_size": 80})

    assert engine_docx.detect_watermark_types(out) == {"text"}

    doc = Document(out)
    dr = _body_marked_drawings(doc)[0]
    assert engine_docx._mark_of(dr) == engine_docx.MARK_TEXT
    # 人为改成衬于文字下方，识别与清除仍必须正常
    dr.find(".//" + qn("wp:anchor")).set("behindDoc", "1")
    doc.save(out)
    assert engine_docx.detect_watermark_types(out) == {"text"}, \
        "识别依赖标记而非 behindDoc：切换层级后仍应识别"
    cleared = os.path.join(d, "cleared.docx")
    core.clear_watermark(out, output_path=cleared, kinds=["text"])
    assert engine_docx.detect_watermark_types(cleared) == set(), \
        "清除依赖标记而非 behindDoc：切换层级后仍应删掉"
    print("[层] behindDoc 改回 1 后仍可识别/删除 —— 标记与层级解耦")


def test_insert_then_delete_roundtrip_with_body_image():
    """带正文图片的文档：插入→删除→重新插入，正文图片必须始终完好。"""
    d = tempfile.mkdtemp()
    src = _make_doc_with_body_image(d)
    wm = os.path.join(d, "wm.docx")
    core.insert_watermark(src, ["text"], output_path=wm,
                          text={"text": "机密", "font_size": 80})

    cleared = os.path.join(d, "cleared.docx")
    core.clear_watermark(wm, output_path=cleared, kinds=["text"])

    doc = Document(cleared)
    body_imgs = [dr for dr in _body_drawings(doc) if engine_docx._mark_of(dr) is None]
    assert len(body_imgs) == 1, "删除水印不得连带删掉正文图片"
    assert "用户正文" in [p.text for p in doc.paragraphs], "正文文字必须保留"

    again = os.path.join(d, "again.docx")
    core.insert_watermark(cleared, ["text"], output_path=again,
                          text={"text": "机密2", "font_size": 80})
    assert engine_docx.detect_watermark_types(again) == {"text"}
    assert len([dr for dr in _body_drawings(Document(again))
                if engine_docx._mark_of(dr) is None]) == 1, "重插后正文图片仍应完好"
    print("[层] 插入→删除→重插 往返：正文图片与文字始终完好")
