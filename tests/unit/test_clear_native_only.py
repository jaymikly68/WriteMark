"""
引擎级测试：验证 `clear_watermark(clear_native_only=True)` 只删「Word 原生/用户自制水印」、
保留本工具添加的水印（含 v1.6.2 起写在正文 body 里的那份）。

这是 v1.6.3 新需求「文档同时含本工具水印与用户水印时，可单独删用户水印」的底层保证。
"""
from __future__ import annotations

import io
import os
import tempfile

from docx import Document
from PIL import Image

from watermark_tool import core, engine_docx


def _png(tmp, name):
    p = os.path.join(tmp, name)
    b = io.BytesIO()
    Image.new("RGB", (10, 10), (1, 1, 1)).save(b, "PNG")
    with open(p, "wb") as f:
        f.write(b.getvalue())
    return p


def _build(tmp, kinds):
    """返回带指定本工具水印 + 1 张「用户 behindDoc 原生图」的 docx 路径。"""
    src = os.path.join(tmp, "in.docx")
    doc = Document()
    doc.add_paragraph("用户正文")
    doc.save(src)
    opts = {}
    if "text" in kinds:
        opts["text"] = {"text": "机密"}
    if "image" in kinds:
        opts["image"] = {"image_path": _png(tmp, "wm.png")}
    wm = os.path.join(tmp, "wm.docx")
    core.insert_watermark(src, list(kinds), output_path=wm, **opts)
    # 用户在页眉里放一张自己的 behindDoc 图（无本工具标记 -> 被识别为原生水印）
    d = Document(wm)
    b = io.BytesIO()
    Image.new("RGB", (10, 10), (2, 2, 2)).save(b, "PNG")
    rid, _ = d.sections[0].header.part.get_or_add_image(io.BytesIO(b.getvalue()))
    engine_docx._add_drawing_to_part(
        d.sections[0].header,
        engine_docx._make_drawing(rid, 300000, 300000, 0, "用户背景图", "", 880, 0, 0))
    d.save(wm)
    return wm


def _native_count(path):
    n = 0
    doc = Document(path)
    for section in doc.sections:
        for part in engine_docx._iter_parts(doc, section, include_footers=True):
            for block in engine_docx._top_graphic_blocks(part._element):
                if engine_docx._is_native_wm_block(block):
                    n += 1
    return n


def _tool_count(path):
    return len(list(engine_docx._iter_marked_drawings(Document(path))))


def test_clear_native_only_removes_native_keeps_tool_text():
    tmp = tempfile.mkdtemp()
    wm = _build(tmp, ["text"])
    assert _native_count(wm) == 1, "前置：应有 1 张用户原生图"
    before = _tool_count(wm)
    assert before >= 1, "前置：应有本工具水印"

    out = os.path.join(tmp, "out.docx")
    core.clear_watermark(wm, output_path=out, clear_native_only=True)

    assert _native_count(out) == 0, "native_only 应删掉用户原生图"
    assert _tool_count(out) == before, "native_only 必须保留本工具水印（含正文 body 那份）"
    assert engine_docx.detect_watermark_types(out) == {"text"}, "本工具文字水印应在"
    print("[engine] clear_native_only=True -> 删用户图、留本工具文字水印")


def test_clear_native_only_removes_native_keeps_tool_image():
    tmp = tempfile.mkdtemp()
    wm = _build(tmp, ["image"])
    assert _native_count(wm) == 1
    before = _tool_count(wm)
    assert before >= 1

    out = os.path.join(tmp, "out.docx")
    core.clear_watermark(wm, output_path=out, clear_native_only=True)

    assert _native_count(out) == 0, "native_only 应删掉用户原生图"
    assert _tool_count(out) == before, "native_only 必须保留本工具图片水印"
    assert engine_docx.detect_watermark_types(out) == {"image"}
    print("[engine] clear_native_only=True -> 删用户图、留本工具图片水印")


def test_clear_native_only_does_not_touch_body_watermark():
    """v1.6.2 水印写在正文 body：native_only 绝不可误删 body 里的本工具标记。"""
    tmp = tempfile.mkdtemp()
    wm = _build(tmp, ["text", "image"])
    body_before = sum(
        1 for _ in engine_docx._iter_marked_drawings(Document(wm))
        if _.getroottree().getpath(_).startswith("/w:document"))  # 粗略：body 内的标记
    # 直接用更可靠的判定：native_only 前后本工具标记总数不变即可证明没误删
    tool_before = _tool_count(wm)
    out = os.path.join(tmp, "out.docx")
    core.clear_watermark(wm, output_path=out, clear_native_only=True)
    assert _tool_count(out) == tool_before, "native_only 不应减少任何本工具标记"
    _ = body_before
    print("[engine] clear_native_only=True 不动正文 body 里的本工具水印")


def test_full_clear_still_removes_both():
    """回归：默认 clear_watermark()（kinds=None）仍同时清掉本工具与原生水印。"""
    tmp = tempfile.mkdtemp()
    wm = _build(tmp, ["text"])
    out = os.path.join(tmp, "out.docx")
    core.clear_watermark(wm, output_path=out)  # 默认全清
    assert _native_count(out) == 0, "默认全清应删用户原生图"
    assert _tool_count(out) == 0, "默认全清应删本工具水印"
    assert engine_docx.detect_watermark_types(out) == set()
    print("[engine] 默认 clear_watermark() -> 本工具与原生一并清除（回归）")


def test_kinds_text_removes_tool_text_keeps_native():
    """kinds=['text'] 仅删本工具文字水印，用户的原生图原样保留（既有行为）。"""
    tmp = tempfile.mkdtemp()
    wm = _build(tmp, ["text"])
    out = os.path.join(tmp, "out.docx")
    core.clear_watermark(wm, output_path=out, kinds=["text"])
    assert engine_docx.detect_watermark_types(out) == set(), "本工具文字水印应清除"
    assert _native_count(out) == 1, "kinds 模式不应碰用户原生图"
    print("[engine] kinds=['text'] -> 删本工具文字、留用户原生图")
