"""
防去除加固测试：平铺满页、多份冗余嵌入、伪装显示名、清除与份数校验。

v1.6.2 起水印改放 **正文 body**（而非页眉/页脚）：浮于正文之上、覆盖每一页。
本文件相应把“扫描页眉/页脚”改为“扫描正文 body”，并把份数期望从「按 part 算」
改成「按正文段落算」（每个段落都会挂一份，重叠成视觉上的一份）。核心不变式
仍是：伪装显示名不命中按名匹配、私有标记藏在 descr、一键清除能把全部水印清干净。
"""
from __future__ import annotations

import os
import re
import tempfile
import zipfile

from docx import Document

sys_path = os.path.dirname(os.path.abspath(__file__))
import sys
sys.path.insert(0, sys_path)

from watermark_tool import core, engine_docx


def _make_docx(path):
    doc = Document()
    doc.add_paragraph("第一段 hello world 测试文档")
    doc.add_paragraph("第二段内容，用于验证水印插入不影响正文。")
    doc.save(path)


def _text_opts(**over):
    d = {"text": "机密 CONFIDENTIAL", "font_size": 120, "color": (192, 0, 0),
         "angle": 30, "transparency": 0.6, "scale": 1.0,
         "offset_x": 0.0, "offset_y": 0.0,
         "cn_font_name": "仿宋", "latin_font_name": "Times New Roman"}
    d.update(over)
    return {"text": d}


def _body_xml(path):
    """返回 word/document.xml 的内容（正文 body）。"""
    z = zipfile.ZipFile(path)
    return [z.read(n).decode("utf8") for n in z.namelist()
            if n == "word/document.xml"]


def _header_footer_xml(path):
    z = zipfile.ZipFile(path)
    hs = [z.read(n).decode("utf8") for n in z.namelist()
          if re.match(r"word/header\d+\.xml$", n)]
    fs = [z.read(n).decode("utf8") for n in z.namelist()
          if re.match(r"word/footer\d+\.xml$", n)]
    return hs, fs


def _names(xml_list):
    out = []
    for x in xml_list:
        out += re.findall(r'<wp:docPr[^>]*name="([^"]+)"', x)
    return out


def _descrs(xml_list):
    out = []
    for x in xml_list:
        out += re.findall(r'<wp:docPr[^>]*descr="([^"]+)"', x)
    return out


def _positions(xml_list):
    """用 lxml 解析出各水印图形的绝对定位坐标集合 {(x, y), ...}（EMU 字符串）。"""
    from lxml import etree
    WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
    out = set()
    for x in xml_list:
        root = etree.fromstring(x.encode("utf8"))
        for anchor in root.iter("{%s}anchor" % WP):
            h = anchor.find(".//{%s}positionH/{%s}posOffset" % (WP, WP))
            v = anchor.find(".//{%s}positionV/{%s}posOffset" % (WP, WP))
            if h is not None and v is not None:
                out.add((h.text, v.text))
    return out


def _n_paragraphs(path):
    """正文段落数（水印按段落数重复，覆盖每一页）。"""
    return len(Document(path).paragraphs)


def test_standard_mode_unchanged():
    """默认（不开加固）：水印落正文、名字就是标记名、不写页眉/页脚。"""
    d = tempfile.mkdtemp(prefix="_hard_")
    p = os.path.join(d, "a.docx")
    _make_docx(p)
    n = _n_paragraphs(p)
    res = engine_docx.insert_watermark(p, ["text"], **_text_opts())
    assert res["inserted"] == n, res            # 每个正文段落一份
    body = _body_xml(p)
    assert _names(body) == [engine_docx.MARK_TEXT] * n, _names(body)
    hs, fs = _header_footer_xml(p)
    assert not any(_names(hs)) and not any(_names(fs)), "标准模式不应写入页眉/页脚"
    print(f"[标准] 正文 {n} 份 + 原名，与历史行为一致（落点改为正文）")


def test_tile_produces_grid():
    """平铺：每个正文段落内写入 rows×cols 份，且分散到不同坐标。"""
    d = tempfile.mkdtemp(prefix="_hard_")
    p = os.path.join(d, "b.docx")
    _make_docx(p)
    n = _n_paragraphs(p)
    res = engine_docx.insert_watermark(p, ["text"], tile=True,
                                       tile_rows=4, tile_cols=3, **_text_opts())
    assert res["inserted"] == n * 12, res
    assert res["tile"] is True
    body = _body_xml(p)
    assert len(_names(body)) == n * 12, len(_names(body))
    # 平铺必须真的分散到不同坐标（不是 12 份叠在正中）
    pos = _positions(body)
    assert len(pos) == 12, f"平铺 12 份应有 12 个不同坐标，实际 {len(pos)}"
    xs = {p[0] for p in pos}
    ys = {p[1] for p in pos}
    assert len(xs) == 3 and len(ys) == 4, f"应为 3 列 × 4 行，实际 {len(xs)}×{len(ys)}"
    print(f"[平铺] 每段 4×3 = 12 份（共 {n*12}），分布在 {len(xs)} 列 × {len(ys)} 行的 {len(pos)} 个坐标")


def test_redundant_uses_decoy_names_but_no_footer():
    """冗余嵌入：显示名伪装（无 watermark 字样）、标记藏在 descr；v1.6.2 起不写页脚。"""
    d = tempfile.mkdtemp(prefix="_hard_")
    p = os.path.join(d, "c.docx")
    _make_docx(p)
    n = _n_paragraphs(p)
    res = engine_docx.insert_watermark(p, ["text"], redundant=True, **_text_opts())
    assert res["inserted"] == n, res            # 冗余只影响显示名，不改变份数
    body = _body_xml(p)
    assert len(_names(body)) == n
    all_names = _names(body)
    assert all("watermark" not in nm.lower() for nm in all_names), all_names
    # 私有标记仍可在 descr 中读到（保证本工具能识别与清除）
    all_descrs = _descrs(body)
    assert all(dc == engine_docx.MARK_TEXT for dc in all_descrs), all_descrs
    hs, fs = _header_footer_xml(p)
    assert not any(_names(hs)) and not any(_names(fs)), "v1.6.2 起不再写页眉/页脚"
    print(f"[冗余] 正文 {n} 份；显示名 {all_names[:3]}（无 watermark 字样），descr 保留标记")


def test_names_not_watermark_like():
    """伪装名不应命中“按名字匹配水印”的常见规则。"""
    d = tempfile.mkdtemp(prefix="_hard_")
    p = os.path.join(d, "d.docx")
    _make_docx(p)
    n = _n_paragraphs(p)
    engine_docx.insert_watermark(p, ["text"], tile=True, tile_rows=3, tile_cols=2,
                                 redundant=True, **_text_opts())
    body = _body_xml(p)
    names = _names(body)
    assert len(names) == n * 6, len(names)     # 每段 6 份
    for nm in names:
        low = nm.lower()
        for bad in ("watermark", "水印", "wbg_watermark", "powerplus"):
            assert bad not in low, f"{nm} 命中关键字 {bad}"
    print(f"[伪装] {len(names)} 份图形名均无 watermark/水印 等关键字；样例 {names[:3]}")


def test_clear_removes_tiled_watermark():
    """一键清除必须能清掉正文里平铺的全部水印（v1.6.2 起水印在正文）。"""
    d = tempfile.mkdtemp(prefix="_hard_")
    p = os.path.join(d, "e.docx")
    _make_docx(p)
    n = _n_paragraphs(p)
    engine_docx.insert_watermark(p, ["text"], tile=True, tile_rows=4, tile_cols=2,
                                 redundant=True, **_text_opts())
    before = engine_docx.count_watermarks(p)
    assert before == n * 8, before
    res = engine_docx.clear_watermark(p)
    assert res["removed"] == n * 8, res
    assert engine_docx.count_watermarks(p) == 0
    assert not engine_docx.has_watermark(p)
    body = _body_xml(p)
    assert not _names(body)
    hs, fs = _header_footer_xml(p)
    assert not _names(hs) and not _names(fs)
    print(f"[清除] 正文平铺共 {before} 份一次清空，正文/页眉/页脚无残留")


def test_count_watermarks():
    d = tempfile.mkdtemp(prefix="_hard_")
    p = os.path.join(d, "f.docx")
    _make_docx(p)
    n = _n_paragraphs(p)
    assert engine_docx.count_watermarks(p) == 0
    engine_docx.insert_watermark(p, ["text"], **_text_opts())
    assert engine_docx.count_watermarks(p) == n
    engine_docx.insert_watermark(p, ["text"], tile=True, tile_rows=5, tile_cols=3,
                                 redundant=True, **_text_opts())
    assert engine_docx.count_watermarks(p) == n * 15
    print(f"[计数] 份数统计正确（0 → {n} → {n*15}）")
