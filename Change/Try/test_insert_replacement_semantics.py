"""钉死 `insert_watermark()` 的**替换语义**（不是追加语义）。

`engine_docx.insert_watermark(path, kinds)` 对每个页眉/页脚 part 先执行
`_remove_in_part(part)`（清掉本工具**所有**已标记的水印图形），再按本次传入的
kinds 重新插入。因此：

    先 insert(["text"])，再单独 insert(["image"])  →  文字水印会消失，只剩图片水印

这不是 bug，而是“每次插入 = 按当前勾选重新生成一份水印”的设计。它的价值在于让
“只勾选图片水印再点一次”能真正去掉文字水印；代价是**如果调用方只传部分 kinds**，
未勾选的那一类会被静默移除。

⚠ 调用姿势提醒（本文件踩过一次）：`core.insert_watermark` 在「输出路径 == 源文件」
时会经 `_protect_original()` 把输出改写为 `xxxWaterMark.docx`，因此要验证
“在原文件上再插一次”的替换行为，必须直接调 `engine_docx.insert_watermark`。

本文件做两件事：
1. 把替换语义用测试固定下来——将来若有人改成“追加”、或反过来以为它是追加而加补丁，
   都会被这几条用例拦住；
2. 确认 GUI 侧每次都传“按勾选得到的完整集合”，即不存在“只传部分 kinds”的调用点。
"""
from __future__ import annotations

import inspect
import os

import pytest
from docx import Document

from watermark_tool import engine_docx

USER_EDIT = "USER_EDIT_12345"


def _png_bytes(color=(12, 34, 56)):
    import io as _io
    from PIL import Image
    buf = _io.BytesIO()
    Image.new("RGBA", (64, 64), color + (255,)).save(buf, "PNG")
    return buf.getvalue()


def _make_source(path):
    """建立一个带正文 + 表格 + 页眉的 .docx。"""
    doc = Document()
    doc.add_paragraph("原始正文第一段")
    doc.add_paragraph("原始正文第二段")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "表A1"
    table.cell(0, 1).text = "表A2"
    table.cell(1, 0).text = "表B1"
    table.cell(1, 1).text = "表B2"
    doc.sections[0].header.paragraphs[0].text = "页眉文字"
    doc.save(path)
    return path


def _text_opts():
    return {"text": {"text": "机密", "font_size": 120, "color": (0, 0, 0),
                     "transparency": 0.5}}


def _img_opts(path):
    img_path = os.path.join(path, "wm.png")
    with open(img_path, "wb") as f:
        f.write(_png_bytes())
    return {"image": {"image_path": img_path, "transparency": 0.5, "scale": 0.3}}


def _types(path):
    return engine_docx.detect_watermark_types(path)


def _body_texts(path):
    return [p.text for p in Document(path).paragraphs]


def _table_grid(path):
    doc = Document(path)
    return [[[c.text for c in row.cells] for row in t.rows] for t in doc.tables]


def _reference_text_only_count(tmp_dir):
    """给出「只插文字」时应有的份数，作为替换后份数的参照基准。"""
    p = os.path.join(tmp_dir, "only_text_ref.docx")
    _make_source(p)
    engine_docx.insert_watermark(p, ["text"], **_text_opts())
    return engine_docx._count_marks_in_doc(Document(p))


# ---------------------------------------------------------------------------
# 1. 替换语义本身
# ---------------------------------------------------------------------------
def test_insert_text_then_image_only_removes_text_watermark(tmp_path):
    """先插 ["text"]，再单独插 ["image"]：文字水印必须消失（替换，不是追加）。"""
    p = _make_source(str(tmp_path / "s.docx"))
    engine_docx.insert_watermark(p, ["text"], **_text_opts())
    assert _types(p) == {"text"}, f"前置条件：应只有文字水印，实际 {_types(p)}"

    engine_docx.insert_watermark(p, ["image"], **_img_opts(str(tmp_path)))
    assert _types(p) == {"image"}, (
        f"替换语义：单独插入图片水印后应只剩图片水印，实际 {_types(p)}")

    print("[insert] 替换语义：文字水印 -> 再插图片 -> 文字水印消失，只剩图片")


def test_insert_both_then_text_only_removes_image_watermark(tmp_path):
    """先插 ["text","image"]，再单独插 ["text"]：图片水印必须消失。"""
    p = _make_source(str(tmp_path / "s.docx"))
    engine_docx.insert_watermark(p, ["text", "image"],
                                 **_text_opts(), **_img_opts(str(tmp_path)))
    assert _types(p) == {"text", "image"}, f"前置条件：两类都应在，实际 {_types(p)}"

    engine_docx.insert_watermark(p, ["text"], **_text_opts())
    assert _types(p) == {"text"}, (
        f"替换语义：只插文字后应只剩文字水印，实际 {_types(p)}")

    print("[insert] 替换语义：两类 -> 只插文字 -> 图片水印消失")


def test_replacement_is_idempotent_not_accumulating(tmp_path):
    """替换不等于“先清再插多留几份”：反复重插后的份数应与单次插入完全一致。"""
    p = _make_source(str(tmp_path / "s.docx"))
    tmp = str(tmp_path)

    engine_docx.insert_watermark(p, ["text", "image"],
                                 **_text_opts(), **_img_opts(tmp))
    both = engine_docx._count_marks_in_doc(Document(p))

    engine_docx.insert_watermark(p, ["text"], **_text_opts())
    after_text = engine_docx._count_marks_in_doc(Document(p))
    reference = _reference_text_only_count(tmp)
    assert after_text == reference, (
        f"替换后份数应与「只插文字」一致（参考 {reference}），实际 {after_text}")

    engine_docx.insert_watermark(p, ["text", "image"],
                                 **_text_opts(), **_img_opts(tmp))
    after_both = engine_docx._count_marks_in_doc(Document(p))
    assert after_both == both, (
        f"再次全量插入应回到 {both} 份，实际 {after_both}"
        f"（说明清理不彻底或层数在累加）")

    print(f"[insert] 替换语义份数守恒：全量={both}，替换后={after_text} -> {after_both}")


def test_replacement_does_not_touch_body_or_table(tmp_path):
    """替换只作用在水印层：正文、表格必须原样。"""
    p = _make_source(str(tmp_path / "s.docx"))
    tmp = str(tmp_path)

    engine_docx.insert_watermark(p, ["text"], **_text_opts())
    doc = Document(p)
    doc.add_paragraph(USER_EDIT)
    doc.save(p)
    body_before, grid_before = _body_texts(p), _table_grid(p)

    engine_docx.insert_watermark(p, ["image"], **_img_opts(tmp))
    assert USER_EDIT in _body_texts(p), "替换插入不得吞掉用户正文"

    engine_docx.insert_watermark(p, ["text"], **_text_opts())
    assert USER_EDIT in _body_texts(p), "再次替换插入不得吞掉用户正文"

    engine_docx.insert_watermark(p, ["text", "image"],
                                 **_text_opts(), **_img_opts(tmp))
    assert USER_EDIT in _body_texts(p), "全量重插后正文仍须在"
    assert _body_texts(p) == body_before, "正文被改写"
    assert _table_grid(p) == grid_before, "表格被改写"

    print("[insert] 替换插入不影响正文/表格：USER_EDIT 三轮重插后仍在")


# ---------------------------------------------------------------------------
# 2. GUI 侧是否始终传完整集合
# ---------------------------------------------------------------------------
def test_gui_gather_kinds_derives_full_set_from_both_toggles():
    """GUI 的 kinds 由「文字水印/图片水印」两个勾选框整体决定，不是增量传参。

    因此 GUI 上不存在“只传部分 kinds 就插入”的调用点，替换语义表现为
    “按当前勾选重新生成一份水印”，不会悄悄删掉用户想要的那一类。
    """
    from watermark_tool import gui as gui_mod

    src_code = inspect.getsource(gui_mod.App._gather_kinds)
    assert "self.text_enabled" in src_code, (
        "_gather_kinds 应依据「是否启用文字水印」给出完整集合")
    assert "self.image_enabled" in src_code, (
        "_gather_kinds 应依据「是否启用图片水印」给出完整集合")
    assert 'kinds.append("text")' in src_code and 'kinds.append("image")' in src_code, (
        "_gather_kinds 应把两种类型都append进 kinds 再返回（完整集合）")
    print("[gui] _gather_kinds 由两个勾选框给出完整 kinds 集合，不存在“只传部分”的调用点")


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__, "-q", "-p", "no:faulthandler"]))
