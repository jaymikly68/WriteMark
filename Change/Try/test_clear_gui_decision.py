"""
GUI「一键清除水印」(`App._clear()`) 的**真实行为**测试。

不看源码文本，而是真的调用 `_clear()`，捕获它交给 `core.clear_watermark` 的调用，
必要时把捕获到的任务**真的执行一遍**，再检查最终文档状态。

要钉死的两条安全契约（都是修复前的真实缺陷）：

1. 按检测到的类型精确删除
   `{text}` -> kinds=['text']；`{image}` -> kinds=['image']。
   修复前 `types` 不等于 `{"text","image"}` 时 kinds 保持 None，
   掉进 `clear_any=True`，把页眉里用户自己的 behindDoc 图也一起删了。

2. 检测失败 -> 停止清除 + 报告，**绝不**退化成扩大删除范围
   修复前 `except Exception: kinds = None`，同样掉进 `clear_any=True`。

说明：`App` 的构造（整窗 UI + 定时器）对这类决策测试是纯开销，
因此这里用 `App.__new__(App)` 造一个不带 UI 的实例、只填 `_clear()` 真正读到的
几个属性。被调用的仍是**真实的 `App._clear()` 方法本身**。
"""
from __future__ import annotations

import os
import sys
import tempfile

import pytest
from docx import Document

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from PySide6.QtWidgets import QLineEdit

from watermark_tool import core, engine_docx
from watermark_tool import gui as gui_mod

USER_BODY = "用户正文"


# ---------------------------------------------------------------------------
# 测试替身
# ---------------------------------------------------------------------------
_WARNINGS: list[tuple[str, str]] = []


class _FakeMessageBox:
    """替身：只记录调用了什么提示框，不发真弹窗。

    注意：不能用实例方法来记，否则 `instance.warning(tuple)` 会又解析到这个
    替身自身、无限递归。这里统一写进模块级列表。
    """

    @staticmethod
    def warning(parent, title, text, *a, **k):
        _WARNINGS.append((title, str(text)))

    @staticmethod
    def information(parent, title, text, *a, **k):
        _WARNINGS.append((title, str(text)))


class _Recorder:
    """替身：记录 core.clear_watermark 的真实调用参数。"""

    def __init__(self):
        self.calls = []

    def __call__(self, src_path, output_path=None, kinds=None, **kw):
        self.calls.append({"src": src_path, "out": output_path, "kinds": kinds})
        return {"ok": True, "engine": "docx", "removed": 0,
                "kinds": sorted(kinds or [])}


def _stub_app(monkeypatch, file_path, out_path):
    """造一个只带 `_clear()` 所需属性的 App 实例，并拦下任务执行。"""
    app = gui_mod.App.__new__(gui_mod.App)
    app.file_path = file_path
    app.out_edit = QLineEdit(out_path)

    jobs = []
    monkeypatch.setattr(gui_mod.App, "_run", lambda self, job: jobs.append(job))
    return app, jobs


def _apply(app, jobs, monkeypatch):
    """执行 `_clear()` 排队下来的清除任务，返回替身记录到的调用参数。

    跑完**立刻 undo 所有 patch**：测试后面还要用真正的 `core.clear_watermark`
    做端到端验证，否则那次调用也会打到替身上、什么文件都不会生成。
    """
    rec = _Recorder()
    monkeypatch.setattr(core, "clear_watermark", rec)
    try:
        for job in jobs:
            job()
    finally:
        monkeypatch.undo()
    return rec.calls


# ---------------------------------------------------------------------------
# 素材
# ---------------------------------------------------------------------------
def _svg_png(tmp, name):
    from PIL import Image
    import io as _io
    b = _io.BytesIO()
    Image.new("RGB", (24, 24), (9, 9, 9)).save(b, "PNG")
    p = os.path.join(tmp, name)
    with open(p, "wb") as f:
        f.write(b.getvalue())
    return p


def _build(tmp, name, kinds, with_user_behind=False):
    """返回 (file_path, output_path)：一份带指定水印的 docx + 一个输出路径。"""
    src = os.path.join(tmp, f"{name}_in.docx")
    doc = Document()
    doc.add_paragraph(USER_BODY)
    doc.save(src)

    opts = {"text": {"text": "机密"},
            "image": {"image_path": _svg_png(tmp, f"{name}.png")}}
    wm = os.path.join(tmp, f"{name}_wm.docx")
    core.insert_watermark(src, list(kinds), output_path=wm, **opts)
    if with_user_behind:
        # 用户在页眉里放一张自己的 behindDoc 图——修复前会被 GUI 的删除一起带走
        d = Document(wm)
        rid, _ = d.sections[0].header.part.get_or_add_image(
            __import__("io").BytesIO(open(_svg_png(tmp, f"{name}b.png"), "rb").read()))
        engine_docx._add_drawing_to_part(
            d.sections[0].header,
            engine_docx._make_drawing(rid, 300000, 300000, 0, "用户背景图", "", 880, 0, 0))
        d.save(wm)
    return wm, os.path.join(tmp, f"{name}_out.docx")


def _user_behind_count(path):
    n = 0
    for dr in Document(path).sections[0].header._element.iter(
            engine_docx.qn("w:drawing")):
        anchor = dr.find(".//" + engine_docx.qn("wp:anchor"))
        if anchor is not None and anchor.get("behindDoc") == "1" \
                and engine_docx._mark_of(dr) is None:
            n += 1
    return n


# ===========================================================================
# P0：只有一种水印时，GUI 必须精确按类型删除
# ===========================================================================
def test_gui_only_text_watermark_clears_text_only(monkeypatch, app):
    """只有文字水印：GUI 必须传 kinds=['text']（而不是 None 的全清）。"""
    tmp = tempfile.mkdtemp()
    fp, out = _build(tmp, "only_text", ["text"], with_user_behind=True)

    app, jobs = _stub_app(monkeypatch, fp, out)
    gui_mod.App._clear(app)
    calls = _apply(app, jobs, monkeypatch)

    assert len(calls) == 1, f"应当恰好发起一次清除，实际 {len(calls)}"
    assert calls[0]["kinds"] == ["text"], (
        f"只有文字水印时应传 kinds=['text']，实际 {calls[0]['kinds']}——"
        "None 会走 clear_any，连带删掉用户自己的 behindDoc 图")

    # 真的执行一遍，看最终文档
    core.clear_watermark(fp, output_path=out, kinds=["text"])
    assert engine_docx.detect_watermark_types(out) == set(), "文字水印应已清除"
    assert _user_behind_count(out) == 1, "用户自己的 behindDoc 图必须保留"
    assert USER_BODY in [p.text for p in Document(out).paragraphs]
    print("[GUI] 只有文字水印 -> kinds=['text']；水印清除，用户 behindDoc 图保留")


def test_gui_only_image_watermark_clears_image_only(monkeypatch, app):
    """只有图片水印：GUI 必须传 kinds=['image']。"""
    tmp = tempfile.mkdtemp()
    fp, out = _build(tmp, "only_img", ["image"], with_user_behind=True)

    app, jobs = _stub_app(monkeypatch, fp, out)
    gui_mod.App._clear(app)
    calls = _apply(app, jobs, monkeypatch)

    assert len(calls) == 1
    assert calls[0]["kinds"] == ["image"], (
        f"只有图片水印时应传 kinds=['image']，实际 {calls[0]['kinds']}")

    core.clear_watermark(fp, output_path=out, kinds=["image"])
    assert engine_docx.detect_watermark_types(out) == set(), "图片水印应已清除"
    assert _user_behind_count(out) == 1, "用户自己的 behindDoc 图必须保留"
    print("[GUI] 只有图片水印 -> kinds=['image']；水印清除，用户 behindDoc 图保留")


# ===========================================================================
# P0：同时含两类水印时弹窗，三个选项分别对应三种 kinds
# ===========================================================================
@pytest.mark.parametrize("choice,expected", [
    ("text", ["text"]),
    ("image", ["image"]),
    ("both", ["text", "image"]),
])
def test_gui_dual_watermark_dialog_maps_choice_to_kinds(monkeypatch, app, choice, expected):
    """双水印文档 -> 弹窗 -> 用户选什么就删什么。"""
    tmp = tempfile.mkdtemp()
    fp, out = _build(tmp, "dual", ["text", "image"])

    app, jobs = _stub_app(monkeypatch, fp, out)
    asked = []
    monkeypatch.setattr(gui_mod.App, "_ask_clear_choice",
                        lambda self: asked.append(True) or choice)
    gui_mod.App._clear(app)
    calls = _apply(app, jobs, monkeypatch)

    assert asked == [True], "双水印文档应该弹一次选择框"
    assert len(calls) == 1 and calls[0]["kinds"] == expected, (
        f"用户选 {choice} 时应传 {expected}，实际 {calls[0] if calls else None}")

    core.clear_watermark(fp, output_path=out, kinds=expected)
    want = set() if choice == "both" else {"text", "image"} - {choice}
    assert engine_docx.detect_watermark_types(out) == want, (
        f"选 {choice} 后应剩 {want or '无'}，实际 {engine_docx.detect_watermark_types(out)}")
    print(f"[GUI] 双水印弹窗：选{choice} -> kinds={expected} -> 残留 {want or '无'}")


def test_gui_dual_watermark_user_cancel_does_nothing(monkeypatch, app):
    """用户点“取消”时，一次清除都不许发生。"""
    tmp = tempfile.mkdtemp()
    fp, out = _build(tmp, "dual_cancel", ["text", "image"])

    app, jobs = _stub_app(monkeypatch, fp, out)
    monkeypatch.setattr(gui_mod.App, "_ask_clear_choice", lambda self: None)
    gui_mod.App._clear(app)

    assert jobs == [], "用户取消时不应排队任何清除任务"
    assert not os.path.exists(out), "用户取消时不该产生输出文件"
    print("[GUI] 双水印弹窗点取消 -> 不发起任何清除")


# ===========================================================================
# P0：检测失败 -> 停止清除、报告原因，不扩大删除范围
# ===========================================================================
def test_gui_detection_failure_does_not_expand_clear_scope(monkeypatch, app):
    """检测抛异常：绝不能退化成 kinds=None（clear_any）。

    同时验证：用户看到了提示，且异常原因被带出来（不吞异常）。
    """
    tmp = tempfile.mkdtemp()
    fp, out = _build(tmp, "fail", ["text"], with_user_behind=True)
    assert _user_behind_count(fp) == 1

    boom = RuntimeError("模拟检测崩溃：文档结构异常")
    monkeypatch.setattr(engine_docx, "detect_watermark_types", lambda p: (_ for _ in ()).throw(boom))

    _WARNINGS.clear()
    monkeypatch.setattr(gui_mod, "QMessageBox", _FakeMessageBox)

    app, jobs = _stub_app(monkeypatch, fp, out)
    gui_mod.App._clear(app)
    calls = _apply(app, jobs, monkeypatch)

    assert calls == [], f"检测失败时不应发起任何清除，实际 {calls}"
    assert not os.path.exists(out), "检测失败时不应产生输出文件"
    assert len(_WARNINGS) == 1, "必须弹窗告知用户"
    title, text = _WARNINGS[0]
    assert "取消" in text and "检测" in text, f"提示语应说明检测失败并已取消，实际：{text}"
    assert "模拟检测崩溃" in text, f"真实原因应透传给用户，实际：{text}"
    print("[GUI] 检测失败 -> 不清除、不扩大范围、弹窗报告原因")


def test_gui_detection_failure_leaves_document_untouched(monkeypatch, app):
    """检测失败后原文件必须一字未改（含用户自己的 behindDoc 图）。"""
    tmp = tempfile.mkdtemp()
    fp, out = _build(tmp, "fail2", ["image"], with_user_behind=True)

    real_detect = engine_docx.detect_watermark_types   # 断言时要用真的（下面会被替换）
    monkeypatch.setattr(engine_docx, "detect_watermark_types",
                        lambda p: (_ for _ in ()).throw(OSError("读取失败")))
    _WARNINGS.clear()
    monkeypatch.setattr(gui_mod, "QMessageBox", _FakeMessageBox)

    app, jobs = _stub_app(monkeypatch, fp, out)
    gui_mod.App._clear(app)
    _apply(app, jobs, monkeypatch)

    assert _user_behind_count(fp) == 1, "原文件的用户 behindDoc 图不应被动过"
    assert real_detect(fp) == {"image"}, "原文件水印应原样"
    print("[GUI] 检测失败后原文档零改动")


# ===========================================================================
# 空集：没有识别到本工具水印时的既有产品语义
# ===========================================================================
def test_gui_no_watermark_detected_keeps_full_clear_semantics(monkeypatch, app):
    """检测为空集：保持 kinds=None（「一键清除水印」连 Word 原生水印一起清）。

    这是产品决策（按钮语义就是“去掉水印”），不是缺陷；但它确实会清掉页眉里
    用户自己的 behindDoc 图，属于已知取舍，故此处**只固化当前产品语义**，
    不把它当成安全保证。
    """
    tmp = tempfile.mkdtemp()
    src = os.path.join(tmp, "plain.docx")
    Document().add_paragraph(USER_BODY)
    doc = Document()
    doc.add_paragraph(USER_BODY)
    doc.save(src)
    fp, out = src, os.path.join(tmp, "plain_out.docx")

    app, jobs = _stub_app(monkeypatch, fp, out)
    gui_mod.App._clear(app)
    calls = _apply(app, jobs, monkeypatch)

    assert len(calls) == 1 and calls[0]["kinds"] is None, (
        f"未检测到本工具水印时应保持 kinds=None（全清语义），实际 {calls}")
    print("[GUI] 未检测到水印 -> kinds=None（保持“连原生水印一起清”的既有语义）")
