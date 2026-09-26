"""
本轮回归的核心文件：把 WriteMark 的「删除文字水印 / 删除图片水印 / 删除全部水印」
三个入口，以及 watchdog 对这三种删除的响应，逐条做成真实回归测试。

与既有测试的区别：
- test_watchdog_preserve.py 关心的是“用户改了正文后，补回是否覆盖”；
- 本文件关心的是**按类型区分**的清除语义——删文字时图片水印必须活着，
  删图片时文字水印必须活着，删全部时正文/表格/用户自己的图片与形状必须活着。

识别方式（按当前代码实现，不是猜的）：
- 文字水印 = 图形上 wp:docPr/@descr（兼容 @name）等于 MARK_TEXT 的浮动图形；
- 图片水印 = 同上等于 MARK_IMG；二者都由 engine_docx._mark_of() 读出。
- engine_docx.detect_watermark_types() 正是用这两个常量区分「有哪些类型」。

三个删除入口（按当前实际调用链）：
- GUI「一键清除水印」 -> gui._clear() -> core.clear_watermark(src, out, kinds=...)
  - kinds=["text"]            -> engine_docx.clear_watermark 的 only_kinds 分支，只删 MARK_TEXT
  - kinds=["image"]           -> 只删 MARK_IMG
  - kinds=["text","image"]    -> 只删这两类标记（**不是** clear_any）
  - kinds=None（文档只有一种类型时会这样传）-> clear_any=True 全清
- watchdog 检测到水印减少 -> watchdog._reinsert() -> core.insert_watermark(..., preserve_existing=True)
  - 复用已有 output 时的“先清”用的是 **kinds 限定清除**（core.py 里显式传 kinds=list(kinds)），
    因此只删文字水印时，图片水印不会被连带清掉。
"""
from __future__ import annotations

import io
import os
import re
import sys
import tempfile

from docx import Document
from docx.shared import Pt
from PIL import Image

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from watermark_tool import core, engine_docx, watchdog as wd_mod

USER_EDIT = "USER_EDIT_12345"


# ---------------------------------------------------------------------------
# 测试素材：正文 + 表格 + 文字水印 + 图片水印
# ---------------------------------------------------------------------------
def _png_bytes(color=(10, 20, 30), size=(40, 40)):
    img = Image.new("RGB", size, color)
    bio = io.BytesIO()
    img.save(bio, "PNG")
    return bio.getvalue()


def _png_file(tmp, name="wm.png", color=(9, 9, 9)):
    path = os.path.join(tmp, name)
    with open(path, "wb") as f:
        f.write(_png_bytes(color))
    return path


def _make_source(path):
    """正文 + 表格，作为插入水印前的原始素材。"""
    doc = Document()
    doc.add_paragraph("原始正文第一段")
    doc.add_paragraph("原始正文第二段")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "表A1"
    t.cell(0, 1).text = "表A2"
    t.cell(1, 0).text = "表B1"
    t.cell(1, 1).text = "表B2"
    doc.save(path)
    return path


def _both_kinds_opts(tmp):
    """同时启用文字水印与图片水印的参数。"""
    return {
        "text": {"text": "机密 CONFIDENTIAL", "font_size": 110, "color": (192, 0, 0),
                 "angle": 30, "transparency": 0.6, "scale": 1.0,
                 "offset_x": 0.0, "offset_y": 0.0},
        "image": {"image_path": _png_file(tmp, "stamp.png", (0, 90, 200)),
                  "angle": 30, "transparency": 0.5, "scale": 1.0,
                  "offset_x": 0.0, "offset_y": 0.0},
    }


def _build_both_watermarks(tmp):
    """造出「带正文 + 表格 + 文字水印 + 图片水印」的 output，并在正文里留下用户标记。

    返回 (src, out)。out 已经带 USER_EDIT（用户随后会在 Word 里编辑它）。
    """
    src = os.path.join(tmp, "src.docx")
    out = os.path.join(tmp, "out.docx")
    _make_source(src)
    core.insert_watermark(src, ["text", "image"], output_path=out,
                          **_both_kinds_opts(tmp))
    # 模拟“用户在 output 上改了正文”（这正是 watchdog 必须保护的东西）
    doc = Document(out)
    doc.add_paragraph(USER_EDIT)
    doc.tables[0].cell(0, 1).text = "用户改过的单元格"
    doc.save(out)
    return src, out


def _types(path):
    """返回该文档中存在的水印类型集合。"""
    return engine_docx.detect_watermark_types(path)


def _body_texts(path):
    return [p.text for p in Document(path).paragraphs]


def _table_grid(path):
    return [[[c.text for c in row.cells] for row in t.rows]
            for t in Document(path).tables]


def _user_shape_in_header(path):
    """页眉里用户自己插入的形状数量（Word 存成 mc:AlternateContent + wps:wsp，
    uri 是 wordprocessingShape，不是 picture）。

    这是最容易被 clear_any 误删的一类对象，所以每个“全清”场景都要盯着它。
    """
    q = engine_docx.qn
    n = 0
    doc = Document(path)
    for section in doc.sections:
        for part in engine_docx._iter_parts(doc, section, include_footers=True):
            for drawing in part._element.iter(q("w:drawing")):
                gd = drawing.find(".//" + q("a:graphicData"))
                if gd is None:
                    continue
                if gd.get("uri") == "http://schemas.microsoft.com/office/word/2010/wordprocessingShape":
                    n += 1
    return n


WPS_URI = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"


def _add_user_shape(path, name="用户插入的形状"):
    """在页眉里放一个「用户自己画的形状」，并且把它设为衬于文字下方（置于底层）。

    Word 保存 2010+ 形状时用的正是 mc:AlternateContent + w:drawing + wps:wsp，
    它的 a:graphicData/@uri 不是图片命名空间——正是最容易被 clear_any 误删的一类对象。
    """
    from lxml import etree
    mc = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
    doc = Document(path)
    hdr = doc.sections[0].header
    el = hdr._element

    ac = etree.SubElement(el, mc + "AlternateContent")
    choice = etree.SubElement(ac, mc + "Choice")
    choice.set(mc + "Requires", "wps")
    drawing = etree.SubElement(choice, engine_docx.qn("w:drawing"))
    anchor = etree.SubElement(drawing, engine_docx.qn("wp:anchor"))
    anchor.set("behindDoc", "1")
    docPr = etree.SubElement(anchor, engine_docx.qn("wp:docPr"))
    docPr.set("id", "777")
    docPr.set("name", name)
    graphic = etree.SubElement(anchor, engine_docx.qn("a:graphic"))
    gd = etree.SubElement(graphic, engine_docx.qn("a:graphicData"))
    gd.set("uri", WPS_URI)
    etree.SubElement(gd, "{" + WPS_URI + "}wsp")

    doc.save(path)
    return path


def _user_delete_marked(path, marks):
    """模拟「用户在 Word 里手动删除水印」：直接摘掉指定标记的水印图形。

    刻意**不调用** WriteMark 的任何删除 API——UserData 侧的行为才是要验证的。
    marks 为 None 表示两类全删。
    """
    doc = Document(path)
    removed = 0
    for section in doc.sections:
        for part in engine_docx._iter_parts(doc, section, include_footers=True):
            for drawing in list(part._element.iter(engine_docx.qn("w:drawing"))):
                mark = engine_docx._mark_of(drawing)
                if mark is None:
                    continue
                if marks is None or mark in marks:
                    engine_docx._detach_drawing(drawing)
                    removed += 1
    doc.save(path)
    return removed


def _run_watchdog(src, out, kinds, opts, expected, settle=1.2):
    """真实启动 watchdog 线程，让它自己发现水印变化并补回。

    这样“外部修改”才真的经过 _loop -> watermark_count -> _reinsert 全链路，
    而不是测试自己调 _reinsert() 装作触发过。

    expected 必须是「用户删除水印之前」的本工具水印总份数：
    watchdog 的判据就是**总份数**（engine_docx.count_watermarks），
    它并不区分被删掉的是文字类还是图片类——这一点由 A5 单独锁定。
    """
    wd = wd_mod.WatermarkWatchdog(src, out, kinds, opts, interval=0.2,
                                  expected=expected)
    wd.start()
    try:
        wd._stop.wait(settle)      # 4~6 个检测周期，足以让它发现一次删除
    finally:
        wd.stop()
    return wd


# ===========================================================================
# A 组：按当前代码确认识别方式与三个删除入口（防止日后改名后「猜错」）
# ===========================================================================
def test_text_and_image_watermarks_are_distinguished_by_private_marks(tmp_path=None):
    """识别层：文字水印与图片水印靠不同的私有标记区分。"""
    import tempfile as _tf
    tmp = str(tmp_path) if tmp_path is not None else _tf.mkdtemp()
    src = _make_source(os.path.join(tmp, "a.docx"))
    opts = _both_kinds_opts(tmp)
    out = os.path.join(tmp, "a_out.docx")
    core.insert_watermark(src, ["text", "image"], output_path=out, **opts)

    assert engine_docx.MARK_TEXT != engine_docx.MARK_IMG, \
        "两个私有标记必须不同，否则无法按类型区分"
    assert _types(out) == {"text", "image"}, (
        f"同时插入两类水印后应检出两种类型，实际 {_types(out)}")

    # 逐个图形核对标记归属
    q = engine_docx.qn
    doc = Document(out)
    for section in doc.sections:
        for part in engine_docx._iter_parts(doc, section, include_footers=True):
            for drawing in part._element.iter(q("w:drawing")):
                mark = engine_docx._mark_of(drawing)
                assert mark in engine_docx.MARK_NAMES, f"本工具水印必须有私有标记，实际 {mark}"
    print("[A1] 文字/图片水印由 MARK_TEXT / MARK_IMG 两个私有标记区分，detect_watermark_types 正确返回两者")


def test_three_user_clear_choices_map_to_three_distinct_kinds(tmp_path=None):
    """GUI 的三个按钮（去文字 / 去图片 / 都去）必须映射到三种不同的 kinds。

    直接从 gui._clear() 的源码里把映射表抠出来求值——代码改了映射，这里立刻失败。
    """
    from watermark_tool import gui as gui_mod

    src_fn = re.search(r'kinds\s*=\s*(\{.*?\})\[choice\]',
                       __import__("inspect").getsource(gui_mod.App._clear), re.S)
    assert src_fn, "gui._clear 里找不到 kinds 映射表（映射结构可能已改动）"
    mapping = eval(src_fn.group(1))            # noqa: S307 - 测试内固定字面量
    assert set(mapping) == {"text", "image", "both"}, \
        f"三个按钮的返回键应保持 text/image/both，实际 {sorted(mapping)}"
    assert mapping["text"] == ["text"], "「去除文字水印」只能传 kinds=['text']"
    assert mapping["image"] == ["image"], "「去除图片水印」只能传 kinds=['image']"
    assert set(mapping["both"]) == {"text", "image"}, "「都去除」必须两类都传"

    # 用户只选一种类型时，_clear 不弹窗、直接把 kinds 留成 None（走全清）
    src_code = __import__("inspect").getsource(gui_mod.App._clear)
    assert 'if types == {"text", "image"}:' in src_code, (
        "只有文档同时含两类水印时才弹窗；否则 kinds=None 走全清（clear_any），"
        "这条行为本测试依赖它，被改动时应同时改这里")
    print(f"[A2] GUI 三按钮 -> kinds 映射：{mapping}")


def test_three_clear_entries_remove_different_watermark_sets(tmp_path=None):
    """三个删除入口（只删文字 / 只删图片 / 全删）实际删掉的东西必须不同。

    注意 core.clear_watermark 的语义是「不破坏原文件」：它先把 src 复制成 output，
    再对副本执行清除。为了让差异可见，src 本身就应当已经带两类水印。
    """
    import tempfile as _tf
    tmp = str(tmp_path) if tmp_path is not None else _tf.mkdtemp()
    opts = _both_kinds_opts(tmp)
    plain = os.path.join(tmp, "c_plain.docx")
    _make_source(plain)
    # 带两类水印的“源”（core.clear_watermark 的输入必须是它，否则看不出差异）
    src = os.path.join(tmp, "c_src.docx")
    core.insert_watermark(plain, ["text", "image"], output_path=src, **opts)
    assert _types(src) == {"text", "image"}

    results = {}
    for tag, kinds in (("text_only", ["text"]), ("image_only", ["image"]),
                       ("both", ["text", "image"])):
        out = os.path.join(tmp, f"c_{tag}.docx")
        res = core.clear_watermark(src, output_path=out, kinds=kinds)
        assert res.get("ok"), f"{tag} 入口调用失败"
        results[tag] = _types(out)

    assert results["text_only"] == {"image"}, (
        f"只删文字水印后应只剩图片水印，实际 {results['text_only']}")
    assert results["image_only"] == {"text"}, (
        f"只删图片水印后应只剩文字水印，实际 {results['image_only']}")
    assert results["both"] == set(), (
        f"删全部后不应有任何本工具水印，实际 {results['both']}")
    print(f"[A3] 只删文字->{{image}}；只删图片->{{text}}；删全部->{{}}")


def test_core_clear_entry_rebuilds_output_from_source(tmp_path=None):
    """文档化 core.clear_watermark 的“不破坏原文件”语义对 output 的影响。

    core.clear_watermark(src, output_path=out) 会先把 src 复制成 out 再清除，
    因此 **out 上已有的用户修改会被 src 覆盖**，这与 core.insert_watermark
    的 preserve_existing 语义不一致。本用例把该行为钉死，防止日后无意改变；
    同时提醒：GUI 的“一键清除”用的正是这条路径（见 D8 技术债务）。
    """
    import tempfile as _tf
    tmp = str(tmp_path) if tmp_path is not None else _tf.mkdtemp()
    plain = _make_source(os.path.join(tmp, "d_plain.docx"))
    opts = {"text": {"text": "X"}}
    src = os.path.join(tmp, "d_src.docx")
    core.insert_watermark(plain, ["text"], output_path=src, **opts)
    assert _types(src) == {"text"}

    out = os.path.join(tmp, "d_out.docx")
    core.insert_watermark(src, ["text"], output_path=out, **opts)
    doc = Document(out)
    doc.add_paragraph(USER_EDIT)
    doc.save(out)
    assert USER_EDIT in _body_texts(out)

    core.clear_watermark(src, output_path=out, kinds=["text"])

    assert USER_EDIT not in _body_texts(out), (
        "core.clear_watermark 从 src 重建 output，用户的编辑本就会被覆盖——"
        "这正是要被文档化的行为，改动它属于产品决策。")
    print("[A4] core.clear_watermark 从 src 重建 output（不破坏原文件语义），用户编辑被覆盖")


# ===========================================================================
# B 组：watchdog 回归三场景（用户明确要求分别验证，不合并）
# ===========================================================================
def test_scenario_delete_text_watermark_only(tmp_path=None):
    """场景 1：只删除文字水印。

    初始：文字水印=有 / 图片水印=有 / USER_EDIT=有
    删除后：文字=无 / 图片=有 / USER_EDIT=有
    触发 watchdog 后：文字=有 / 图片=有 / USER_EDIT=仍在
    """
    import tempfile as _tf
    tmp = str(tmp_path) if tmp_path is not None else _tf.mkdtemp()
    src, out = _build_both_watermarks(tmp)
    opts = _both_kinds_opts(tmp)

    expected = core.watermark_count(out)
    assert _types(out) == {"text", "image"}, "前置条件：两类水印都在"
    assert expected >= 2, f"前置条件：文字+图片应算两份水印，实际 {expected}"
    assert USER_EDIT in _body_texts(out)

    # --- 用户选择「删除文字水印」---
    engine_docx.clear_watermark(out, kinds=["text"])
    assert _types(out) == {"image"}, "前置条件：应只剩图片水印"
    assert USER_EDIT in _body_texts(out), "删除水印这一步不该动正文"

    # --- 触发 watchdog（真实线程）---
    _run_watchdog(src, out, ["text", "image"], opts, expected)

    after = _types(out)
    assert after == {"text", "image"}, (
        f"只删文字水印后，watchdog 应把文字水印补回来，且图片水印仍在，实际 {after}")
    assert USER_EDIT in _body_texts(out), "watchdog 补回后用户新增的正文必须还在"
    assert "原始正文第二段" in _body_texts(out), "原正文必须还在"
    assert _table_grid(out)[0][0] == ["表A1", "用户改过的单元格"], "表格内容必须没变"
    print(f"[B1] 只删文字水印（{expected}->{expected-1} 份）-> watchdog 补回文字水印，"
          f"图片水印与 USER_EDIT 均保留")


def test_scenario_delete_image_watermark_only(tmp_path=None):
    """场景 2：只删除图片水印（与场景 1 镜像，必须单独验证）。"""
    import tempfile as _tf
    tmp = str(tmp_path) if tmp_path is not None else _tf.mkdtemp()
    src, out = _build_both_watermarks(tmp)
    opts = _both_kinds_opts(tmp)

    expected = core.watermark_count(out)
    assert _types(out) == {"text", "image"}

    # --- 用户选择「删除图片水印」---
    engine_docx.clear_watermark(out, kinds=["image"])
    assert _types(out) == {"text"}, "前置条件：应只剩文字水印"
    assert USER_EDIT in _body_texts(out)

    _run_watchdog(src, out, ["text", "image"], opts, expected)

    after = _types(out)
    assert after == {"text", "image"}, (
        f"只删图片水印后，watchdog 应把图片水印补回来，且文字水印仍在，实际 {after}")
    assert USER_EDIT in _body_texts(out), "watchdog 补回后用户新增的正文必须还在"
    assert _table_grid(out)[0][0] == ["表A1", "用户改过的单元格"], "表格内容必须没变"
    print(f"[B2] 只删图片水印（{expected}->{expected-1} 份）-> watchdog 补回图片水印，"
          f"文字水印与 USER_EDIT 均保留")


def test_scenario_delete_all_watermarks(tmp_path=None):
    """场景 3：删除全部水印。正文/表格/USER_EDIT 必须毫发无伤。"""
    import tempfile as _tf
    tmp = str(tmp_path) if tmp_path is not None else _tf.mkdtemp()
    src, out = _build_both_watermarks(tmp)
    opts = _both_kinds_opts(tmp)

    grid_before = _table_grid(out)
    body_before = _body_texts(out)
    expected = core.watermark_count(out)

    # --- 用户选择「文字和图片水印都去除」---
    engine_docx.clear_watermark(out, kinds=["text", "image"])
    assert _types(out) == set(), f"前置条件：两类水印都应被删，实际 {_types(out)}"
    assert USER_EDIT in _body_texts(out)

    _run_watchdog(src, out, ["text", "image"], opts, expected)

    assert _types(out) == {"text", "image"}, f"删全部后 watchdog 应把两类都补回，实际 {_types(out)}"
    assert _body_texts(out) == body_before, (
        "watchdog 补回后正文发生变化——用户修改被覆盖了。\n"
        f"before={body_before}\n\nafter={_body_texts(out)}")
    assert _table_grid(out) == grid_before, "表格内容必须没变"
    print("[B3] 删全部 -> watchdog 补回两类水印，正文/表格/USER_EDIT 全部原样")


# ===========================================================================
# C 组：模拟“用户在 Word 里直接删除”（不调用 WriteMark 的删除 API）
# ===========================================================================
def test_user_manually_deletes_text_watermark_via_xml(tmp_path=None):
    """用户在 Word 里手动删掉文字水印 -> watchdog 必须补文字、留图片、留正文。"""
    import tempfile as _tf
    tmp = str(tmp_path) if tmp_path is not None else _tf.mkdtemp()
    src, out = _build_both_watermarks(tmp)
    opts = _both_kinds_opts(tmp)
    expected = core.watermark_count(out)

    removed = _user_delete_marked(out, {engine_docx.MARK_TEXT})
    assert removed > 0, "前置条件：应摘到文字水印图形"
    assert _types(out) == {"image"}, f"删掉文字水印后应只剩图片水印，实际 {_types(out)}"

    _run_watchdog(src, out, ["text", "image"], opts, expected)

    assert _types(out) == {"text", "image"}, \
        f"用户手动删文字水印后 watchdog 应补回文字水印且保留图片水印，实际 {_types(out)}"
    assert USER_EDIT in _body_texts(out)
    print("[C1] 用户手动删文字水印 -> watchdog 只补文字，图片水印与正文不受影响")


def test_user_manually_deletes_image_watermark_via_xml(tmp_path=None):
    """用户在 Word 里手动删掉图片水印 -> watchdog 必须补图片、留文字、留正文。"""
    import tempfile as _tf
    tmp = str(tmp_path) if tmp_path is not None else _tf.mkdtemp()
    src, out = _build_both_watermarks(tmp)
    opts = _both_kinds_opts(tmp)
    expected = core.watermark_count(out)

    removed = _user_delete_marked(out, {engine_docx.MARK_IMG})
    assert removed > 0, "前置条件：应摘到图片水印图形"
    assert _types(out) == {"text"}, f"删掉图片水印后应只剩文字水印，实际 {_types(out)}"

    _run_watchdog(src, out, ["text", "image"], opts, expected)

    assert _types(out) == {"text", "image"}, \
        f"用户手动删图片水印后 watchdog 应补回图片水印且保留文字水印，实际 {_types(out)}"
    assert USER_EDIT in _body_texts(out)
    print("[C2] 用户手动删图片水印 -> watchdog 只补图片，文字水印与正文不受影响")


def test_user_manually_deletes_both_watermarks_via_xml(tmp_path=None):
    """用户在 Word 里把两类水印都删掉 -> watchdog 两块都要补回来。"""
    import tempfile as _tf
    tmp = str(tmp_path) if tmp_path is not None else _tf.mkdtemp()
    src, out = _build_both_watermarks(tmp)
    opts = _both_kinds_opts(tmp)
    expected = core.watermark_count(out)

    removed = _user_delete_marked(out, None)
    assert removed >= 2, f"前置条件：两类水印都应被摘掉，实际 {removed}"
    assert _types(out) == set()

    _run_watchdog(src, out, ["text", "image"], opts, expected)

    assert _types(out) == {"text", "image"}, \
        f"用户在 Word 里删掉全部水印后，watchdog 应把两类都补回，实际 {_types(out)}"
    assert USER_EDIT in _body_texts(out)
    assert "原始正文第一段" in _body_texts(out)
    print("[C3] 用户手动删两类水印 -> watchdog 全量补回，正文与 USER_EDIT 保留")


# ===========================================================================
# D 组：clear_watermark / clear_any 的误删边界（watchdog 补回依赖它）
# ===========================================================================
def test_clear_text_kind_does_not_touch_image_watermark_or_user_objects(tmp_path=None):
    """只删文字水印时，图片水印 + 正文/表格/页眉页脚里用户自己的东西都不能少。"""
    import tempfile as _tf
    tmp = str(tmp_path) if tmp_path is not None else _tf.mkdtemp()
    src, out = _build_both_watermarks(tmp)
    _add_user_shape(out)

    body_before, grid_before = _body_texts(out), _table_grid(out)
    shape_before = _user_shape_in_header(out)
    assert shape_before == 1

    # 这正是 watchdog 补回时 core.py 调用的那个函数（就地清除，不从 src 覆盖）
    engine_docx.clear_watermark(out, kinds=["text"])

    assert _types(out) == {"image"}, f"应只剩图片水印，实际 {_types(out)}"
    assert _body_texts(out) == body_before, "只删文字水印时正文不该有任何变化"
    assert _table_grid(out) == grid_before, "只删文字水印时表格不该有任何变化"
    assert _user_shape_in_header(out) == shape_before, "用户页眉形状不该被动"
    print("[D1] 只删文字水印：图片水印保留，正文/表格/用户页眉形状零变化")


def test_clear_image_kind_does_not_touch_text_watermark_or_user_objects(tmp_path=None):
    """只删图片水印时，文字水印 + 正文/表格/用户对象都要原样。"""
    import tempfile as _tf
    tmp = str(tmp_path) if tmp_path is not None else _tf.mkdtemp()
    src, out = _build_both_watermarks(tmp)
    _add_user_shape(out)

    body_before, grid_before = _body_texts(out), _table_grid(out)
    shape_before = _user_shape_in_header(out)

    engine_docx.clear_watermark(out, kinds=["image"])

    assert _types(out) == {"text"}, f"应只剩文字水印，实际 {_types(out)}"
    assert _body_texts(out) == body_before, "只删图片水印时正文不该有任何变化"
    assert _table_grid(out) == grid_before, "只删图片水印时表格不该有任何变化"
    assert _user_shape_in_header(out) == shape_before, "用户页眉形状不该被动"
    print("[D2] 只删图片水印：文字水印保留，正文/表格/用户页眉形状零变化")


def test_clear_all_preserves_body_table_user_picture_shape_and_header_footer(tmp_path=None):
    """删全部水印（clear_any 路径）时，正文/表格/用户正文图/用户形状/页眉页脚不能丢。

    注：用户放在页眉里、且**衬于文字下方**的 behindDoc 图片，在这一路径下会被清掉
    ——这是 clear_any 的设计（见 engine_docx.clear_watermark 的 docstring），
    本用例因此**不放置**这类对象，避免把「设计行为」误当成「缺陷」钉死。
    """
    import tempfile as _tf
    tmp = str(tmp_path) if tmp_path is not None else _tf.mkdtemp()
    src, out = _build_both_watermarks(tmp)
    _add_user_shape(out)

    # 用户自己的正文图片（行内）、页眉行内图、页脚文字——都不该被删
    doc = Document(out)
    doc.add_picture(io.BytesIO(_png_bytes((200, 30, 30), (30, 30))), width=Pt(80))
    doc.sections[0].header.paragraphs[0].add_run().add_picture(
        io.BytesIO(_png_bytes((7, 7, 7))))
    doc.sections[0].footer.paragraphs[0].text = "页脚：内部资料，不是水印"
    doc.save(out)

    # 两种“全清”入口都走一遍：
    #  kinds=["text","image"] = GUI 弹窗里点“都去除”
    #  kinds=None              = 文档只有一种类型时 gui._clear 传的默认值（clear_any）
    for label, kinds in (("kinds=text+image", ["text", "image"]), ("kinds=None(clear_any)", None)):
        src2, out2 = _build_both_watermarks(tmp)
        _add_user_shape(out2)
        doc = Document(out2)
        doc.add_picture(io.BytesIO(_png_bytes((200, 30, 30), (30, 30))), width=Pt(80))
        doc.sections[0].header.paragraphs[0].add_run().add_picture(
            io.BytesIO(_png_bytes((7, 7, 7))))
        doc.sections[0].footer.paragraphs[0].text = "页脚：内部资料，不是水印"
        doc.save(out2)

        b0, g0 = _body_texts(out2), _table_grid(out2)
        s0 = _user_shape_in_header(out2)
        d0 = Document(out2)
        i0 = len(d0.inline_shapes)
        q = engine_docx.qn
        hi0 = sum(1 for _ in d0.sections[0].header._element.iter(q("wp:inline")))
        f0 = d0.sections[0].footer.paragraphs[0].text

        engine_docx.clear_watermark(out2, kinds=kinds)

        assert _types(out2) == set(), \
            f"[{label}] 全清后不应剩本工具水印，实际 {_types(out2)}"
        assert _body_texts(out2) == b0, f"[{label}] 全清时正文必须原样"
        assert _table_grid(out2) == g0, f"[{label}] 全清时表格必须原样"
        assert _user_shape_in_header(out2) == s0, \
            f"[{label}] 全清时用户自己的页眉形状被误删了"
        d1 = Document(out2)
        assert len(d1.inline_shapes) == i0, f"[{label}] 用户插入的正文图片被误删"
        assert sum(1 for _ in d1.sections[0].header._element.iter(q("wp:inline"))) == hi0, \
            f"[{label}] 页眉里的普通图片被误删"
        assert d1.sections[0].footer.paragraphs[0].text == f0, \
            f"[{label}] 页脚文字被误删"

    print("[D3] 删全部水印（两种入口）：正文/表格/正文图片/页眉普通图/页脚文字/用户形状 全部保留")


def test_watchdog_scenarios_do_not_touch_user_assets(tmp_path=None):
    """把三个 watchdog 场景跑一遍，逐一确认用户资产（表格/形状/页眉页脚）零变化。"""
    import tempfile as _tf
    tmp = str(tmp_path) if tmp_path is not None else _tf.mkdtemp()
    opts = _both_kinds_opts(tmp)

    # kinds 取 output 里实际存在的两类（用户启用两类守护时就是这个值）
    for tag, marks in (("只删文字", {engine_docx.MARK_TEXT}),
                       ("只删图片", {engine_docx.MARK_IMG}),
                       ("删全部", None)):
        src, out = _build_both_watermarks(tmp)
        _add_user_shape(out)
        before = {
            "body": _body_texts(out),
            "grid": _table_grid(out),
            "shape": _user_shape_in_header(out),
        }
        expected = core.watermark_count(out)

        _user_delete_marked(out, marks)
        _run_watchdog(src, out, ["text", "image"], opts, expected)

        assert _types(out) == {"text", "image"}, \
            f"[{tag}] watchdog 未把两类水印都补回，实际 {_types(out)}"
        assert _body_texts(out) == before["body"], f"[{tag}] 正文被改动"
        assert _table_grid(out) == before["grid"], f"[{tag}] 表格被改动"
        assert _user_shape_in_header(out) == before["shape"], f"[{tag}] 用户页眉形状被误删"
        assert USER_EDIT in _body_texts(out), f"[{tag}] USER_EDIT 丢失"
    print("[D4] watchdog 三场景：正文/表格/用户页眉形状/USER_EDIT 全部零变化")


# ===========================================================================
# E 组：反向验证——证明上面这些测试真能抓住 84a773e 修复的问题
# ===========================================================================
def test_legacy_reinsert_without_preserve_existing_loses_user_edit(tmp_path=None):
    """把 core.insert_watermark 强制退回「84a773e 之前」的行为（从 source 覆盖重建），
    上面那些“USER_EDIT 必须还在”的断言就应当**失败**——这正是本套测试的杀伤力所在。

    旧版 watchdog._reinsert 的源码是：
        core.insert_watermark(src, kinds, output_path=out, **opts)   # 无 preserve_existing
    而 core.insert_watermark 默认从 source copy2 到 output，用户修改被整份覆盖。

    这里用 monkeypatch 复刻该行为，并断言「用户编辑确实会丢失」。
    该断言成立 == 新行为确实修复了真实问题；若将来 preserve_existing 被改坏，
    本用例会失败，提醒守门。
    """
    import tempfile as _tf
    from unittest import mock
    tmp = str(tmp_path) if tmp_path is not None else _tf.mkdtemp()
    src, out = _build_both_watermarks(tmp)
    opts = _both_kinds_opts(tmp)
    assert USER_EDIT in _body_texts(out)

    real_insert = core.insert_watermark

    def legacy_insert(*args, **kwargs):
        # 84a773e 之前没有 preserve_existing：一律从 source 覆盖 output
        kwargs["preserve_existing"] = False
        return real_insert(*args, **kwargs)

    wd = wd_mod.WatermarkWatchdog(src, out, ["text", "image"], opts, interval=0.2)
    wd.expected = core.watermark_count(out)

    with mock.patch.object(core, "insert_watermark", legacy_insert):
        assert _user_delete_marked(out, None) >= 2
        wd._reinsert()                      # 走的是“旧”代码路径

    legacy_texts = _body_texts(out)
    assert USER_EDIT not in legacy_texts, (
        "旧行为（从 source 覆盖）本应把用户的 USER_EDIT 抹掉；"
        "如果这条断言失败，说明 preserve_existing 已失效，"
        "那么 B/C 组那些“USER_EDIT 必须还在”的回归测试也就不复存在保护力了。")

    # 同样的场景，现行行为必须保住 USER_EDIT（正面对照）
    src2, out2 = _build_both_watermarks(tmp)
    assert USER_EDIT in _body_texts(out2)
    wd2 = wd_mod.WatermarkWatchdog(src2, out2, ["text", "image"], opts, interval=0.2)
    wd2.expected = core.watermark_count(out2)
    _user_delete_marked(out2, None)
    wd2._reinsert()
    assert USER_EDIT in _body_texts(out2), "现行行为必须保住用户编辑"
    assert _types(out2) == {"text", "image"}, "现行行为必须同时补回两类水印"
    print("[E1] 反向验证：旧行为丢失 USER_EDIT，新行为（preserve_existing=True）保住它")
