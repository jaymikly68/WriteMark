"""
Word 删除功能的**安全边界**回归：按类型删除时，绝不能扩大删除范围。

本文件锁定的产品契约（WriteMark 的「一键清除水印」）：
    1. 只删文字水印 -> 图片水印必须活着；
    2. 只删图片水印 -> 文字水印必须活着；
    3. 两者都删    -> 两类都没了，但**用户自己的内容**一件都不能少；
    4. 检测失败    -> 停止本次清除并报告，绝不退化成“扩大删除范围”。

与 test_clear_kinds_matrix.py 的分工：
- 后者考察“三个删除入口分别删掉哪些本工具水印”；
- 本文件考察“用户原有的东西有没有被误删”——这才是安全边界。

两套 API 的“不破坏原文件”语义（python-docx 层与 core 层一致）：
- insert：output 与 src 同名时 core 会自动改名，因此本文件一律写到**新文件**；
- clear ：core.clear_watermark 把 src 复制成 output 再清 output，
  **断言必须落在 output 上**，src 保持原样。

已知的设计取舍（不在本文件钉死，见 test_clear_any_scope.py 的 xfail）：
`clear_watermark(kinds=None)` 走 clear_any，会连页眉里**用户自己的**
behindDoc 浮动图片一起删掉。那是「去掉原本就带水印的 Word」这一产品需求的
已知代价。但**按类型删除（kinds 给定）绝不能走这条路**——本文件多处断言
「kinds 限定路径下用户 behindDoc 图片必须活着」，正是为了防止它溜进来。
"""
from __future__ import annotations

import io
import os
import sys
import tempfile

import pytest
from docx import Document
from PIL import Image

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from watermark_tool import core, engine_docx

USER_BODY = "这是用户自己的正文，删除水印后必须还在"


# ---------------------------------------------------------------------------
# 素材与工具
# ---------------------------------------------------------------------------
def _png_bytes(color=(11, 22, 33), size=(32, 32)):
    b = io.BytesIO()
    Image.new("RGB", size, color).save(b, "PNG")
    return b.getvalue()


def _insert_opts(tmp, tag):
    img = os.path.join(tmp, f"{tag}_stamp.png")
    with open(img, "wb") as f:
        f.write(_png_bytes())
    return {
        "text": {"text": "机密文件", "font_size": 100, "angle": 30},
        "image": {"image_path": img},
    }


def _plain_doc(path):
    doc = Document()
    doc.add_paragraph(USER_BODY)
    doc.add_paragraph("第二段正文")
    doc.save(path)
    return path


def _watermarked(tmp, name, kinds):
    """造一份「正文 + 指定类型水印」的文档，返回其路径。

    刻意与 core.insert_watermark 的 output 保护机制协作：写到新文件，
    不覆盖、也不触发 core 的 “output 与 src 同名就改名” 分支。
    """
    src = os.path.join(tmp, f"{name}_src.docx")
    _plain_doc(src)
    out = os.path.join(tmp, f"{name}_wm.docx")
    core.insert_watermark(src, list(kinds), output_path=out,
                          **_insert_opts(tmp, name))
    return out


def _clear(src, kinds=None):
    """执行清除，返回 (result, output_path)。断言请落在 output_path 上。"""
    out = os.path.join(os.path.dirname(src),
                       f"{os.path.basename(src)[:-len('.docx')]}_clr.docx")
    res = core.clear_watermark(src, output_path=out, kinds=kinds)
    return res, out


def _types(path):
    return engine_docx.detect_watermark_types(path)


def _body_texts(path):
    return [p.text for p in Document(path).paragraphs]


def _header_elements(path):
    """所有页眉的 lxml 元素（默认页眉；有则一并含首页/偶数页）。"""
    doc = Document(path)
    out = [doc.sections[0].header._element]
    for attr in ("first_page_header", "even_page_header"):
        try:
            el = getattr(doc.sections[0], attr)._element
            if el is not None and len(el) and el not in out:
                out.append(el)
        except Exception:
            pass
    return out


def _behind(drawing):
    anchor = drawing.find(".//" + engine_docx.qn("wp:anchor"))
    return anchor is not None and anchor.get("behindDoc") == "1"


def _user_behind_count(path):
    """页眉里「behindDoc=1 且无本工具标记」的浮动图形数量——最容易被误删的一类。"""
    return sum(1 for hdr in _header_elements(path)
               for dr in hdr.iter(engine_docx.qn("w:drawing"))
               if _behind(dr) and engine_docx._mark_of(dr) is None)


def _marked_left(path):
    return sum(1 for hdr in _header_elements(path)
               for dr in hdr.iter(engine_docx.qn("w:drawing"))
               if engine_docx._mark_of(dr) is not None)


def _inline_in_header(path):
    return sum(1 for hdr in _header_elements(path)
               for _ in hdr.iter(engine_docx.qn("wp:inline")))


def _header_texts(path):
    doc = Document(path)
    return [[p.text for p in sec.header.paragraphs] for sec in doc.sections]


def _add_user_behind_picture(path, docpr_id=930):
    """用户在页眉里放一张自己的 behindDoc 浮动图片（形状与 Word 原生水印一致）。"""
    doc = Document(path)
    hdr = doc.sections[0].header
    rid, _ = hdr.part.get_or_add_image(io.BytesIO(_png_bytes((200, 180, 0))))
    engine_docx._add_drawing_to_part(
        hdr, engine_docx._make_drawing(rid, 300000, 300000, 0,
                                       "用户页眉背景图", "", docpr_id, 0, 0))
    doc.save(path)
    return path


def _add_user_inline_logo(path):
    """用户在页眉里放一张行内 Logo（最常见的页眉图片形态）。"""
    doc = Document(path)
    doc.sections[0].header.paragraphs[0].add_run().add_picture(
        io.BytesIO(_png_bytes((250, 250, 250))))
    doc.save(path)
    return path


def _weld_watermark_into_text_run(path, mark, user_text, docpr_id=940):
    """把水印图形**塞进一条已经含用户文字的 w:r** 里。

    模拟的现实结构：Word 重新排版合并 run、或旧版本工具留下的结构，
    都会出现「同一条 run 里 w:t 与 w:drawing 并存」。这时删除水印，
    绝不能把那条 run 整条删掉，否则用户文字陪葬。

    mark 取 engine_docx.MARK_TEXT / MARK_IMG，descr 与之一致（本工具识别依据）。
    """
    doc = Document(path)
    hdr = doc.sections[0].header
    if not hdr.paragraphs:
        hdr.add_paragraph()
    run = hdr.paragraphs[0].add_run(user_text)
    rid, _ = hdr.part.get_or_add_image(io.BytesIO(_png_bytes()))
    run._r.append(engine_docx._make_drawing(rid, 300000, 300000, 0,
                                            mark, mark, docpr_id, 0, 0))
    doc.save(path)
    return path


# ===========================================================================
# P0-1 / P0-2：单类型文档的按类型删除（验证最终文档状态，不只看有没有抛异常）
# ===========================================================================
def test_p0_1_text_only_watermark_clear_text_keeps_body():
    """P0-1：文档只有文字水印 -> kinds=['text'] -> 水印消失、正文完好。"""
    tmp = tempfile.mkdtemp()
    src = _watermarked(tmp, "t1", ["text"])
    assert _types(src) == {"text"}, "前置条件：应只有文字水印"

    _, out = _clear(src, kinds=["text"])

    assert _types(out) == set(), f"文字水印应被清掉，实际残留 {_types(out)}"
    assert USER_BODY in _body_texts(out), "用户正文被误删"
    assert Document(out).paragraphs is not None, "文档应仍可被正常打开"
    print("[P0-1] 只有文字水印 + 删文字 -> 水印消失，正文完好")


def test_p0_2_image_only_watermark_clear_image_keeps_body():
    """P0-2：文档只有图片水印 -> kinds=['image'] -> 水印消失、正文完好。"""
    tmp = tempfile.mkdtemp()
    src = _watermarked(tmp, "t2", ["image"])
    assert _types(src) == {"image"}, "前置条件：应只有图片水印"

    _, out = _clear(src, kinds=["image"])

    assert _types(out) == set(), f"图片水印应被清掉，实际残留 {_types(out)}"
    assert USER_BODY in _body_texts(out), "用户正文被误删"
    print("[P0-2] 只有图片水印 + 删图片 -> 水印消失，正文完好")


def test_p0_3_p0_4_p0_5_clear_matrix_on_dual_watermark():
    """P0-3/4/5：文字+图片共存时，三种删除方式的结果矩阵（独立性才是关键）。"""
    tmp = tempfile.mkdtemp()
    cases = [
        # (tag, 删除方式, 期望残留的水印类型)
        ("to_text", ["text"], {"image"}),        # P0-3 只删文字 -> 图片水印活着
        ("to_image", ["image"], {"text"}),       # P0-4 只删图片 -> 文字水印活着
        ("to_both", ["text", "image"], set()),   # P0-5 全删 -> 一个不剩
    ]
    for tag, kinds, want in cases:
        src = _watermarked(tmp, f"m_{tag}", ["text", "image"])
        assert _types(src) == {"text", "image"}, "前置条件：两类水印都在"

        _, out = _clear(src, kinds=kinds)

        assert _types(out) == want, (
            f"[{tag}] kinds={kinds}：应剩 {want}，实际剩 {_types(out)}")
        assert USER_BODY in _body_texts(out), f"[{tag}] 用户正文被误删"
        assert Document(out).paragraphs is not None, f"[{tag}] 文档应仍可正常打开"
    print("[P0-3/4/5] 双水印矩阵：删文字->剩图片；删图片->剩文字；全删->空；正文始终完好")


# ===========================================================================
# P0-6 / P0-7：用户内容保护（水印与用户内容**共用一条 run** 的最坏情况）
# ===========================================================================
def test_p0_6_user_text_in_same_run_as_text_watermark_survives():
    """P0-6：同一条 w:r 里既有用户文字、又有文字水印。

    删文字水印时必须只摘掉图形，用户文字要留着（_detach_drawing 的回归点）。
    """
    tmp = tempfile.mkdtemp()
    src = _watermarked(tmp, "u6", ["text"])
    _weld_watermark_into_text_run(src, engine_docx.MARK_TEXT, "用户页眉自己的说明文字")

    assert _types(src) == {"text"}
    assert "用户页眉自己的说明文字" in _header_texts(src)[0], "前置条件：同 run 文字应在"

    _, out = _clear(src, kinds=["text"])

    assert _types(out) == set(), f"文字水印应被清掉，实际残留 {_types(out)}"
    assert "用户页眉自己的说明文字" in _header_texts(out)[0], (
        "同 run 内的用户文字被水印的删除操作连带删掉了——"
        "正确做法是只摘掉 w:drawing，run 里还有内容就保留 run。")
    assert USER_BODY in _body_texts(out)
    print("[P0-6] 同 run「用户文字 + 文字水印」-> 只删图形，用户文字保留")


def test_p0_6b_user_text_in_same_run_as_image_watermark_survives():
    """同一条 run 里既有用户文字、又有图片水印时，删图片水印也不能连带删文字。"""
    tmp = tempfile.mkdtemp()
    src = _watermarked(tmp, "u6b", ["image"])
    _weld_watermark_into_text_run(src, engine_docx.MARK_IMG, "用户页眉保留文字B")

    _, out = _clear(src, kinds=["image"])

    assert _types(out) == set(), f"图片水印应被清掉，实际残留 {_types(out)}"
    assert "用户页眉保留文字B" in _header_texts(out)[0], "同 run 用户文字被连带删除"
    print("[P0-6b] 同 run「用户文字 + 图片水印」-> 只删图形，用户文字保留")


def test_p0_7_user_logo_and_behind_picture_survive_image_clear():
    """P0-7：用户 Logo（行内）+ 用户 behindDoc 图 + 图片水印，只删图片水印。

    除了图片水印，另外两样都必须活着。
    """
    tmp = tempfile.mkdtemp()
    src = _watermarked(tmp, "u7", ["image"])
    _add_user_inline_logo(src)
    _add_user_behind_picture(src)

    behind0, inline0 = _user_behind_count(src), _inline_in_header(src)
    assert behind0 == 1 and inline0 == 1, f"前置条件：两种用户图都应在，实际 {behind0}/{inline0}"

    _, out = _clear(src, kinds=["image"])

    assert _types(out) == set(), f"图片水印应被清掉，实际残留 {_types(out)}"
    assert _user_behind_count(out) == behind0, "用户自己的 behindDoc 图片被误删"
    assert _inline_in_header(out) == inline0, "用户页眉 Logo（行内图）被误删"
    assert USER_BODY in _body_texts(out)
    print("[P0-7] 删图片水印：水印消失，用户 Logo 与 behindDoc 图一律保留")


# ===========================================================================
# P0-8：behindDoc 普通图片 + WriteMark 水印，三种删除都不得碰它
# ===========================================================================
def test_p0_8_user_behind_picture_survives_all_three_clear_choices():
    """P0-8：用户 behindDoc 图 + 两类水印；只删文字 / 只删图片 / 全删，用户图都得在。"""
    tmp = tempfile.mkdtemp()
    cases = [("a", ["text"]), ("b", ["image"]), ("c", ["text", "image"])]
    for tag, kinds in cases:
        src = _watermarked(tmp, f"u8_{tag}", ["text", "image"])
        _add_user_behind_picture(src, docpr_id=931)
        before = _user_behind_count(src)
        assert before == 1

        _, out = _clear(src, kinds=kinds)

        assert _user_behind_count(out) == before, (
            f"[{tag}] 用户自己的 behindDoc 图片被 kinds={kinds} 这条路径删除了——"
            "按类型删除只能动本工具标记的水印")
        assert USER_BODY in _body_texts(out)
    print("[P0-8] 用户 behindDoc 图：只删文字 / 只删图片 / 全删 三种路径下均保留")


# ===========================================================================
# 修复点回归：单类型文档 + 用户 behindDoc 图（修复前会把用户的图一起删掉）
# ===========================================================================
@pytest.mark.parametrize("kind", ["text", "image"])
def test_single_type_doc_clear_does_not_eat_user_behind_picture(kind):
    """回归：**只有一种水印**时，GUI 传出的 kinds 不能再是 None（clear_any）。

    修复前 gui._clear() 只在 `types == {"text","image"}` 时设置 kinds，
    单类型时会把 kinds 留成 None -> core.clear_watermark 走 clear_any=True，
    于是页眉里**用户自己的** behindDoc 图会跟着本工具水印一起被删。
    这里用与 GUI 完全等价的决策逻辑把它钉死。
    """
    tmp = tempfile.mkdtemp()
    src = _watermarked(tmp, f"s_{kind}", [kind])
    _add_user_behind_picture(src, docpr_id=932)

    # 复刻修复后 gui._clear() 的决策：单类型 -> kinds=[该类型]
    detected = engine_docx.detect_watermark_types(src)
    if detected == {"text", "image"}:
        kinds = ["text", "image"]
    elif detected == {"text"}:
        kinds = ["text"]
    elif detected == {"image"}:
        kinds = ["image"]
    else:
        kinds = None
    assert kinds == [kind], f"前置条件：单类型应映射为 kinds=[{kind}]，实际 {kinds}"

    _, out = _clear(src, kinds=kinds)

    assert _user_behind_count(out) == 1, (
        f"[{kind}] 单类型删除把用户自己的 behindDoc 图删掉了"
        f"（残留 {_user_behind_count(out)}，应仍为 1）")
    assert _marked_left(out) == 0, f"[{kind}] 本工具水印应被清干净，残留 {_marked_left(out)}"
    print(f"[回归] 单类型({kind})删除：用户 behindDoc 图保留，水印清空")


def test_kinds_limited_path_only_removes_marked_graphics():
    """扫描所有单类型文档：kinds 给定时，带标记的图形一定要掉、无标记的一定要留。"""
    tmp = tempfile.mkdtemp()
    for kind in ("text", "image"):
        src = _watermarked(tmp, f"k_{kind}", [kind])
        _add_user_behind_picture(src, docpr_id=933)
        _add_user_inline_logo(src)

        _, out = _clear(src, kinds=[kind])

        assert _marked_left(out) == 0, f"[{kind}] 带标记的图形应全被删掉"
        assert _user_behind_count(out) == 1, f"[{kind}] 无标记 behindDoc 图被误删"
    print("[边界] kinds 限定路径：只删带标记的图形，无标记用户图形零影响")
