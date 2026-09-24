"""
验证引擎：文本 + 图像水印可同时插入、可分别/整体清除、has_watermark 识别两种标记。
"""
import io
import os
import tempfile

from PIL import Image

from watermark_tool import engine_docx, core


def make_docx(path):
    from docx import Document
    d = Document()
    d.add_paragraph("这是一段用于测试的普通正文，不应被水印操作影响。")
    d.save(path)


def make_img(path):
    img = Image.new("RGBA", (200, 200), (255, 0, 0, 255))
    img.save(path, "PNG")


def count_marks(path):
    from docx import Document
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    from docx.oxml.ns import qn
    d = Document(path)
    text_n = img_n = 0
    for s in d.sections:
        for h in (s.header,):
            for drawing in h._element.iter(qn("w:drawing")):
                docPr = drawing.find(".//" + qn("wp:docPr"))
                if docPr is not None:
                    nm = docPr.get("name")
                    if nm == engine_docx.MARK_TEXT: text_n += 1
                    elif nm == engine_docx.MARK_IMG: img_n += 1
    return text_n, img_n


def main():
    tmp = tempfile.mkdtemp()
    docx_path = os.path.join(tmp, "t.docx")
    img_path = os.path.join(tmp, "wm.png")
    make_docx(docx_path)
    make_img(img_path)

    opts = {
        "text": {
            "text": "机密 CONFIDENTIAL",
            "font_name": "微软雅黑",
            "font_size": 120,
            "color": (128, 128, 128),
            "angle": 30.0,
            "transparency": 0.4,
            "scale": 1.25,
        },
        "image": {
            "image_path": img_path,
            "angle": 60.0,
            "transparency": 0.2,
            "scale": 0.8,
        },
    }

    # 1) 同时插入文本 + 图像
    r = engine_docx.insert_watermark(docx_path, ["text", "image"], **opts)
    print("insert both:", r)
    t, i = count_marks(docx_path)
    print("after insert -> text marks:", t, "image marks:", i)
    assert t >= 1 and i >= 1, "同时插入失败：文本或图像标记缺失"
    assert core.has_watermark(docx_path) is True

    # 2) 只插入文本
    r = engine_docx.insert_watermark(docx_path, ["text"], **opts)
    print("insert text only:", r)
    t, i = count_marks(docx_path)
    print("after text-only -> text marks:", t, "image marks:", i)
    assert t >= 1 and i == 0, "仅文本插入时图像标记应被清掉"

    # 3) 只插入图像
    r = engine_docx.insert_watermark(docx_path, ["image"], **opts)
    print("insert image only:", r)
    t, i = count_marks(docx_path)
    print("after image-only -> text marks:", t, "image marks:", i)
    assert t == 0 and i >= 1, "仅图像插入时文本标记应被清掉"

    # 4) 同时插入后再整体清除
    engine_docx.insert_watermark(docx_path, ["text", "image"], **opts)
    r = engine_docx.clear_watermark(docx_path)
    print("clear:", r)
    t, i = count_marks(docx_path)
    print("after clear -> text marks:", t, "image marks:", i)
    assert t == 0 and i == 0, "清除失败：仍有水印标记残留"
    assert core.has_watermark(docx_path) is False

    print("\n全部断言通过：文本/图像水印可同时添加、可独立切换、可整体清除。")
    # 预览渲染（文本+图像两层）
    img = __import__("watermark_tool.preview", fromlist=["render_preview"]).render_preview(opts, ["text", "image"])
    buf = io.BytesIO(); img.save(buf, "PNG")
    print("预览图字节数:", len(buf.getvalue()))


if __name__ == "__main__":
    main()
